"""Spreadsheet formatting with small OOXML changes, preserving unrelated parts."""
import copy,csv,io,os,tempfile
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PySide6.QtCore import Qt,QSignalBlocker
from PySide6.QtGui import QColor,QFont
from PySide6.QtWidgets import QApplication,QTableWidgetItem,QInputDialog,QColorDialog,QComboBox,QSpinBox
from office_editing import command
from office_surgery import S,_active_sheet


def selected_cells(editor):
    ranges=editor.grid.selectedRanges()
    if not ranges:return [(max(0,editor.grid.currentRow()),max(0,editor.grid.currentColumn()))]
    cells={(r,c) for area in ranges for r in range(area.topRow(),area.bottomRow()+1) for c in range(area.leftColumn(),area.rightColumn()+1)}
    if len(cells)>10000:raise ValueError('Format up to 10,000 cells at a time.')
    return sorted(cells)


def paint_cell(item,changes):
    font=item.font()
    for key,value in changes.items():
        if key=='name':font.setFamily(value)
        elif key=='size':font.setPointSizeF(value)
        elif key=='bold':font.setBold(value)
        elif key=='italic':font.setItalic(value)
        elif key=='underline':font.setUnderline(value)
        elif key=='color':item.setForeground(QColor(value))
        elif key=='fill':item.setBackground(QColor(value))
        elif key=='horizontal':item.setTextAlignment({'left':Qt.AlignLeft,'center':Qt.AlignHCenter,'right':Qt.AlignRight}[value]|Qt.AlignVCenter)
    item.setFont(font)


def format_cells(editor,**changes):
    if 'size' in changes and not 6<=changes['size']<=96:raise ValueError('Font size must be 6–96 points.')
    with QSignalBlocker(editor.grid):
        for row,col in selected_cells(editor):
            item=editor.grid.item(row,col)
            if item is None:item=QTableWidgetItem('');editor.grid.setItem(row,col,item)
            paint_cell(item,changes);editor.pending_styles.setdefault((row,col),{}).update(changes)
    editor.dirty=True


def save_cell_formats(editor,path):
    if not editor.pending_styles:return
    tag=lambda n:'{'+S+'}'+n
    with ZipFile(path) as archive:
        style=etree.fromstring(archive.read('xl/styles.xml'));sheet_path=_active_sheet(archive);sheet=etree.fromstring(archive.read(sheet_path));data=sheet.find(tag('sheetData'))
        fonts=style.find(tag('fonts'));fills=style.find(tag('fills'));xfs=style.find(tag('cellXfs'));numfmts=style.find(tag('numFmts'))
        if numfmts is None:numfmts=etree.Element(tag('numFmts'),count='0');style.insert(0,numfmts)
        def append(parent,node):
            parent.append(node);parent.set('count',str(len(parent)));return len(parent)-1
        for (row,col),changes in editor.pending_styles.items():
            address=f'{get_column_letter(col+1)}{row+1}';rownode=next((n for n in data if n.get('r')==str(row+1)),None)
            if rownode is None:
                rownode=etree.Element(tag('row'),r=str(row+1));data.append(rownode)
            cell=next((n for n in rownode if n.get('r')==address),None)
            if cell is None:cell=etree.SubElement(rownode,tag('c'),r=address)
            xf=copy.deepcopy(xfs[int(cell.get('s','0'))]);font=copy.deepcopy(fonts[int(xf.get('fontId','0'))]);font_changed=False
            for key,name in [('name','name'),('size','sz'),('bold','b'),('italic','i'),('underline','u'),('color','color')]:
                if key not in changes:continue
                font_changed=True
                for old in list(font.findall(tag(name))):font.remove(old)
                value=changes[key]
                if key in ('bold','italic','underline'):
                    if value:etree.SubElement(font,tag(name))
                else:etree.SubElement(font,tag(name),**({'rgb':'FF'+value.lstrip('#').upper()} if key=='color' else {'val':str(value)}))
            if font_changed:xf.set('fontId',str(append(fonts,font)));xf.set('applyFont','1')
            if 'fill' in changes:
                fill=etree.Element(tag('fill'));pattern=etree.SubElement(fill,tag('patternFill'),patternType='solid');etree.SubElement(pattern,tag('fgColor'),rgb='FF'+changes['fill'].lstrip('#').upper());etree.SubElement(pattern,tag('bgColor'),indexed='64');xf.set('fillId',str(append(fills,fill)));xf.set('applyFill','1')
            if 'horizontal' in changes:
                alignment=xf.find(tag('alignment'))
                if alignment is None:alignment=etree.SubElement(xf,tag('alignment'))
                alignment.set('horizontal',changes['horizontal']);xf.set('applyAlignment','1')
            if 'number_format' in changes:
                code=changes['number_format'];existing=next((n for n in numfmts if n.get('formatCode')==code),None)
                if existing is None:
                    number=max([163]+[int(n.get('numFmtId')) for n in numfmts])+1;existing=etree.Element(tag('numFmt'),numFmtId=str(number),formatCode=code);append(numfmts,existing)
                xf.set('numFmtId',existing.get('numFmtId'));xf.set('applyNumberFormat','1')
            cell.set('s',str(append(xfs,xf)))
        # Excel expects row and cell records in coordinate order.
        from openpyxl.utils.cell import coordinate_from_string,column_index_from_string
        for row in data:
            row[:]=sorted(row,key=lambda n:column_index_from_string(coordinate_from_string(n.get('r'))[0]) if n.tag==tag('c') else 100000)
        data[:]=sorted(data,key=lambda n:int(n.get('r','0')))
        with tempfile.NamedTemporaryFile(dir=Path(path).parent,suffix='.xlsx',delete=False) as stream:temporary=Path(stream.name)
        try:
            with ZipFile(temporary,'w') as output:
                for info in archive.infolist():output.writestr(info,etree.tostring(style) if info.filename=='xl/styles.xml' else etree.tostring(sheet) if info.filename==sheet_path else archive.read(info.filename))
        except Exception:temporary.unlink(missing_ok=True);raise
    try:os.replace(temporary,path)
    finally:temporary.unlink(missing_ok=True)
    editor.pending_styles.clear()


def load_cell_formats(editor,path):
    editor.pending_styles={}
    if Path(path).suffix.lower()!='.xlsx':return
    book=load_workbook(path,read_only=True,data_only=False)
    try:
        with QSignalBlocker(editor.grid):
            for row in book.active:
                for cell in row:
                    if not getattr(cell,'has_style',False):continue
                    item=editor.grid.item(cell.row-1,cell.column-1)
                    if item is None:item=QTableWidgetItem('');editor.grid.setItem(cell.row-1,cell.column-1,item)
                    values=dict(name=cell.font.name or 'Calibri',size=cell.font.sz or 11,bold=bool(cell.font.b),italic=bool(cell.font.i),underline=bool(cell.font.u))
                    if cell.font.color and cell.font.color.type=='rgb':values['color']='#'+cell.font.color.rgb[-6:]
                    if cell.fill.patternType=='solid' and cell.fill.fgColor.type=='rgb':values['fill']='#'+cell.fill.fgColor.rgb[-6:]
                    if cell.alignment.horizontal in ('left','center','right'):values['horizontal']=cell.alignment.horizontal
                    paint_cell(item,values)
    finally:book.close()


def paste_cells(editor,text=None):
    text=QApplication.clipboard().text() if text is None else text;rows=list(csv.reader(io.StringIO(text),delimiter='\t'))
    if not rows:return
    if sum(len(row) for row in rows)>10000:raise ValueError('Paste up to 10,000 cells at a time.')
    r,c=max(0,editor.grid.currentRow()),max(0,editor.grid.currentColumn())
    editor.grid.setRowCount(max(editor.grid.rowCount(),r+len(rows)));editor.grid.setColumnCount(max(editor.grid.columnCount(),c+max(map(len,rows))));editor._headers()
    for i,row in enumerate(rows):
        for j,value in enumerate(row):
            item=editor.grid.item(r+i,c+j)
            if item:item.setText(value)
            else:editor.grid.setItem(r+i,c+j,QTableWidgetItem(value))
    editor.dirty=True


def setup_cell_tools(editor):
    editor.pending_styles={}
    def copy_cells():
        ranges=editor.grid.selectedRanges()
        if len(ranges)!=1:raise ValueError('Select one rectangular range.')
        area=ranges[0];stream=io.StringIO();writer=csv.writer(stream,delimiter='\t',lineterminator='\n');writer.writerows([editor.grid.item(r,c).text() if editor.grid.item(r,c) else '' for c in range(area.leftColumn(),area.rightColumn()+1)] for r in range(area.topRow(),area.bottomRow()+1));QApplication.clipboard().setText(stream.getvalue())
    command(editor,'Home','Copy',copy_cells,'copy','Ctrl+C');command(editor,'Home','Paste',lambda:paste_cells(editor),'paste','Ctrl+V')
    def toggle(key):
        item=editor.grid.currentItem();font=item.font() if item else QFont();format_cells(editor,**{key:not getattr(font,{'bold':'bold','italic':'italic','underline':'underline'}[key])()})
    for label,key in [('Bold','bold'),('Italic','italic'),('Underline','underline')]:command(editor,'Home',label,lambda k=key:toggle(k),key)
    def color(key):
        value=QColorDialog.getColor(QColor('#ffdd22'),editor)
        if value.isValid():format_cells(editor,**{key:value.name()})
    command(editor,'Home','Text color',lambda:color('color'),'color');command(editor,'Home','Cell fill',lambda:color('fill'),'color')
    for label in ('Left','Center','Right'):command(editor,'Home','Align '+label.lower(),lambda v=label.lower():format_cells(editor,horizontal=v),label.lower())
    def number_format():
        values={'General':'General','Number':'0.00','Currency':'#,##0.00','Percent':'0.00%','Date':'yyyy-mm-dd'};label,ok=QInputDialog.getItem(editor,'Number format','Format',list(values),0,False)
        if ok:format_cells(editor,number_format=values[label]);editor.statusBar().showMessage('Number format saved in XLSX; the grid shows raw values.')
    command(editor,'Home','Number format',number_format,'calculate')
    def font():
        from PySide6.QtWidgets import QFontDialog
        value,ok=QFontDialog.getFont(editor.grid.currentItem().font() if editor.grid.currentItem() else QFont('Calibri',11),editor)
        if ok:format_cells(editor,name=value.family(),size=value.pointSizeF(),bold=value.bold(),italic=value.italic(),underline=value.underline())
    command(editor,'Home','Font',font,'title')
    command(editor,'View','Auto-fit columns',editor.grid.resizeColumnsToContents,'column')
    def filter_rows():
        text,ok=QInputDialog.getText(editor,'Filter current column','Contains (empty shows all rows):')
        if ok:
            column=max(0,editor.grid.currentColumn())
            for row in range(1,editor.grid.rowCount()):
                item=editor.grid.item(row,column);editor.grid.setRowHidden(row,bool(text) and text.casefold() not in (item.text() if item else '').casefold())
    command(editor,'Data','Filter column',filter_rows,'find')
    command(editor,'Data','Show all rows',lambda:[editor.grid.setRowHidden(r,False) for r in range(editor.grid.rowCount())],'table')
