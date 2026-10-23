from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.pagebreak import Break
from pptx import Presentation
from pptx.util import Inches as PInches
from PIL import Image, ImageDraw

root = Path(__file__).resolve().parent
root.mkdir(exist_ok=True)
image = root / 'layout_image.png'
im = Image.new('RGB', (320, 120), 'white')
draw = ImageDraw.Draw(im)
draw.rectangle((5, 5, 314, 114), outline='navy', width=4)
draw.text((20, 48), 'Image anchor', fill='black')
im.save(image)

doc = Document()
section = doc.sections[0]
section.header.paragraphs[0].text = 'Header: render baseline'
section.footer.paragraphs[0].text = 'Footer: render baseline'
p = doc.add_paragraph('Heading in Cambria')
p.runs[0].font.name = 'Cambria'
p.runs[0].font.size = Pt(20)
doc.add_paragraph('Paragraph in Arial with a DATE field: ')
p = doc.paragraphs[-1]
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'DATE')
p._p.append(field)
table = doc.add_table(rows=3, cols=2)
for i, row in enumerate(table.rows):
    for j, cell in enumerate(row.cells):
        cell.text = f'Cell {i},{j}'
doc.add_picture(str(image), width=Inches(2.2))
p = doc.add_paragraph()
p.add_run().add_break(WD_BREAK.PAGE)
doc.add_paragraph('Explicit second page; page break should survive.')
doc.save(root / 'layout.docx')

wb = Workbook()
ws = wb.active
ws.title = 'Print layout'
for row in [['Item', 'Value'], ['A', 12], ['B', 18], ['C', 9]]:
    ws.append(row)
ws['A6'] = 'Formula'
ws['B6'] = '=SUM(B2:B4)'
chart = BarChart()
chart.title = 'Values'
chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
ws.add_chart(chart, 'D2')
ws.print_area = 'A1:L30'
ws.oddHeader.center.text = 'Header'
ws.oddFooter.center.text = 'Footer'
ws.row_breaks.append(Break(id=15))
wb.save(root / 'layout.xlsx')

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
box = slide.shapes.add_textbox(PInches(1), PInches(1), PInches(7), PInches(1))
box.text_frame.text = 'Slide layout in Arial'
box.text_frame.paragraphs[0].runs[0].font.name = 'Arial'
slide.shapes.add_picture(str(image), PInches(1), PInches(2), width=PInches(3))
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.shapes.add_textbox(PInches(1), PInches(1), PInches(7), PInches(1)).text_frame.text = 'Second slide'
prs.save(root / 'layout.pptx')

(root / 'continuation.fdx').write_text('''<?xml version="1.0" encoding="UTF-8" standalone="no"?>
<FinalDraft DocumentType="Script" Template="No" Version="1"><Content>
<Paragraph Type="Scene Heading"><Text>INT. OFFICE - DAY</Text></Paragraph>
<Paragraph Type="Action"><Text>A page-sensitive action beat.</Text></Paragraph>
<Paragraph Type="Character"><Text>ALEX</Text></Paragraph>
<Paragraph Type="Dialogue"><Text>''' + 'Long dialogue. ' * 900 + '''</Text></Paragraph>
</Content></FinalDraft>''', encoding='utf-8')
print('Created DOCX, XLSX, PPTX, FDX and PNG fixtures in', root)
