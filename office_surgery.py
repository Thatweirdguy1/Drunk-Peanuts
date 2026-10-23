"""Small OOXML edits that leave unrelated package parts untouched."""

from __future__ import annotations

import re
import os
import tempfile
from bisect import bisect_right
from difflib import SequenceMatcher
from itertools import accumulate
from pathlib import Path
from posixpath import join, normpath
from zipfile import ZipFile

from lxml import etree
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter, range_boundaries


S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P = "http://schemas.openxmlformats.org/package/2006/relationships"
XML = "http://www.w3.org/XML/1998/namespace"
PARSER = etree.XMLParser(resolve_entities=False, no_network=True)


def replace_runs(paragraph, value: str) -> None:
    """Keep existing run formatting for unchanged text and nearby insertions."""
    runs = list(paragraph.runs)
    if any(any(etree.QName(child).localname not in {"rPr", "t"} for child in run._r) for run in runs):
        raise ValueError("This text contains objects or fields that cannot be edited safely.")
    old = "".join(run.text for run in runs)
    if old != paragraph.text:
        raise ValueError("This text contains a structure that cannot be mapped safely.")
    if not runs:
        if value:
            paragraph.add_run(value)
        return
    ends = list(accumulate(len(run.text) for run in runs))
    pieces = [""] * len(runs)
    for operation, i1, i2, j1, j2 in SequenceMatcher(None, old, value, autojunk=False).get_opcodes():
        if operation == "equal":
            for offset, char in enumerate(value[j1:j2], i1):
                pieces[min(bisect_right(ends, offset), len(runs) - 1)] += char
        elif operation in {"replace", "insert"}:
            pieces[min(bisect_right(ends, i1), len(runs) - 1)] += value[j1:j2]
    for run, text in zip(runs, pieces):
        run.text = text


def replace_text_frame(shape, value: str) -> None:
    paragraphs = list(shape.text_frame.paragraphs)
    lines = value.split("\n")
    if len(lines) != len(paragraphs):
        raise ValueError("Paragraph count changed in a formatted slide text box.")
    for paragraph, line in zip(paragraphs, lines):
        if paragraph.text != line:
            replace_runs(paragraph, line)


def _active_sheet(archive: ZipFile) -> str:
    workbook = etree.fromstring(archive.read("xl/workbook.xml"), PARSER)
    sheets = workbook.find(f"{{{S}}}sheets")
    view = workbook.find(f"{{{S}}}bookViews/{{{S}}}workbookView")
    selected = int(view.get("activeTab", "0")) if view is not None else 0
    sheet = list(sheets)[selected]
    rel_id = sheet.get(f"{{{R}}}id")
    rels = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"), PARSER)
    target = next((rel.get("Target") for rel in rels if rel.get("Id") == rel_id), None)
    if not target:
        raise ValueError("Cannot locate the active worksheet.")
    path = normpath(join("xl", target)) if not target.startswith("/") else target.lstrip("/")
    if not path.startswith("xl/") or path not in archive.namelist():
        raise ValueError("Invalid worksheet path in workbook.")
    return path


def _cell_value(cell, value: str) -> None:
    for child in list(cell):
        if etree.QName(child).localname in {"f", "v", "is"}:
            cell.remove(child)
    if value.startswith("="):
        if len(value) == 1:
            raise ValueError("A formula needs an expression after =.")
        cell.attrib.pop("t", None)
        etree.SubElement(cell, f"{{{S}}}f").text = value[1:]
    elif re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?(?:[eE][+-]?\d+)?", value):
        cell.set("t", "n")
        etree.SubElement(cell, f"{{{S}}}v").text = value
    elif value:
        cell.set("t", "inlineStr")
        inline = etree.SubElement(cell, f"{{{S}}}is")
        node = etree.SubElement(inline, f"{{{S}}}t")
        if value != value.strip():
            node.set(f"{{{XML}}}space", "preserve")
        node.text = value
    else:
        cell.attrib.pop("t", None)


def edit_xlsx_cells(source: Path, target: Path, before: list[list[str]], after: list[list[str]]) -> None:
    """Edit only changed cells in the active sheet; retain every other OOXML part."""
    with ZipFile(source) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)) or any(name.startswith("_xmlsignatures/") for name in names):
            raise ValueError("Duplicate package parts or signed workbooks cannot be edited safely.")
        sheet_path = _active_sheet(archive)
        root = etree.fromstring(archive.read(sheet_path), PARSER)
        data = root.find(f"{{{S}}}sheetData")
        if data is None:
            raise ValueError("This workbook has no editable worksheet data.")
        existing_rows = {int(row.get("r")): row for row in data.findall(f"{{{S}}}row")}
        merged = [range_boundaries(node.get("ref")) for node in root.findall(f"{{{S}}}mergeCells/{{{S}}}mergeCell")]
        count = max(len(before), len(after))
        for row_no in range(1, count + 1):
            old_row = before[row_no - 1] if row_no <= len(before) else []
            new_row = after[row_no - 1] if row_no <= len(after) else []
            for col_no in range(1, max(len(old_row), len(new_row)) + 1):
                old = old_row[col_no - 1] if col_no <= len(old_row) else ""
                new = new_row[col_no - 1] if col_no <= len(new_row) else ""
                if new == old:
                    continue
                if any(left <= col_no <= right and top <= row_no <= bottom and (col_no, row_no) != (left, top)
                       for left, top, right, bottom in merged):
                    raise ValueError("Cannot edit a non-anchor cell in a merged range.")
                row = existing_rows.get(row_no)
                if row is None:
                    row = etree.Element(f"{{{S}}}row", r=str(row_no))
                    following = next((item for item in data if int(item.get("r", "0")) > row_no), None)
                    data.insert(data.index(following) if following is not None else len(data), row)
                    existing_rows[row_no] = row
                address = f"{get_column_letter(col_no)}{row_no}"
                cell = next((item for item in row if item.tag == f"{{{S}}}c" and item.get("r") == address), None)
                if cell is None:
                    cell = etree.Element(f"{{{S}}}c", r=address)
                    following = next((item for item in row if item.tag == f"{{{S}}}c" and
                                      column_index_from_string(coordinate_from_string(item.get("r"))[0]) > col_no), None)
                    row.insert(row.index(following) if following is not None else len(row), cell)
                formula = cell.find(f"{{{S}}}f")
                if formula is not None and formula.get("t") in {"shared", "array", "dataTable"}:
                    raise ValueError("This formula type cannot be edited safely.")
                _cell_value(cell, new)
        replacement = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
        with tempfile.NamedTemporaryFile(dir=target.parent, suffix=".xlsx", delete=False) as pending:
            output_path = Path(pending.name)
        try:
            with ZipFile(output_path, "w") as output:
                for info in archive.infolist():
                    output.writestr(info, replacement if info.filename == sheet_path else archive.read(info.filename))
        except Exception:
            output_path.unlink(missing_ok=True)
            raise
    try:
        os.replace(output_path, target)
    except Exception:
        output_path.unlink(missing_ok=True)
        raise
