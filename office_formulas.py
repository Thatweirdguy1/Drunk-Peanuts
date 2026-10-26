"""Bounded spreadsheet evaluation and result-only painting; raw cells remain editable."""
import ast,math,operator,re
from decimal import Decimal,ROUND_HALF_UP,InvalidOperation
from openpyxl.formula.tokenizer import Tokenizer,TokenizerError
from openpyxl.utils.cell import coordinate_to_tuple,get_column_letter
from PySide6.QtCore import Qt,QTimer,QSignalBlocker
from PySide6.QtWidgets import QStyledItemDelegate,QTableWidgetItem
from office_editing import command

NUMBER_FORMAT_ROLE=Qt.UserRole+25


class FormulaError(ValueError):
    def __init__(self,code,message):super().__init__(message);self.code=code


def evaluator(rows):
    cache={};visiting=set()
    def coords(name):
        try:
            row,col=coordinate_to_tuple(name.replace('$','').upper())
            if not 1<=row<=1048576 or not 1<=col<=16384:raise ValueError()
            return row-1,col-1
        except (ValueError,TypeError):raise FormulaError('#REF!','Invalid cell reference.')
    def numeric(value):
        if value is None:return 0.0
        if not isinstance(value,(int,float)):raise FormulaError('#VALUE!','Arithmetic requires numbers.')
        result=float(value)
        if not math.isfinite(result):raise FormulaError('#NUM!','Numeric result is out of range.')
        return result
    def expression(raw):
        if len(raw)>2000:raise FormulaError('#NUM!','Formula is too long.')
        parts=[]
        for token in Tokenizer(raw).items:
            kind,sub,value=token.type,token.subtype,token.value
            if kind=='WHITE-SPACE':continue
            if kind=='OPERAND':
                if sub=='RANGE':
                    refs=value.split(':')
                    if not all(re.fullmatch(r'\$?[A-Za-z]{1,3}\$?[1-9]\d*',r) for r in refs) or len(refs)>2:raise FormulaError('#NAME?','Named, cross-sheet and external references are not supported yet.')
                    parts.append(('CELL' if len(refs)==1 else 'RANGE')+'('+','.join(repr(r) for r in refs)+')')
                elif sub=='TEXT':parts.append(repr(value[1:-1].replace('""','"')))
                elif sub=='LOGICAL':parts.append('True' if value.upper()=='TRUE' else 'False')
                elif sub=='NUMBER':parts.append(repr(numeric(float(value))))
                elif sub=='ERROR':parts.append('ERROR('+repr(value)+')')
                else:raise FormulaError('#NAME?','Unsupported operand.')
            elif kind=='FUNC':parts.append(value.upper())
            elif kind=='PAREN':parts.append(value)
            elif kind=='SEP' and sub=='ARG':parts.append(',')
            elif kind in ('OPERATOR-INFIX','OPERATOR-PREFIX') and value in ('+','-','*','/','^','&','=','<>','>','<','>=','<='):
                parts.append({'^':'**','=':'==','<>':'!='}.get(value,value))
            else:raise FormulaError('#NAME?','Unsupported formula syntax.')
        return ast.parse(''.join(parts),mode='eval').body
    def value(name):
        key=coords(name)
        if key in cache:
            if isinstance(cache[key],FormulaError):raise cache[key]
            return cache[key]
        if key in visiting:raise FormulaError('#CYCLE!','Circular formula reference.')
        if len(visiting)>=100 or len(cache)>20000:raise FormulaError('#NUM!','Calculation limit reached.')
        visiting.add(key)
        try:
            r,c=key;raw=rows[r][c] if r<len(rows) and c<len(rows[r]) else ''
            if isinstance(raw,str) and raw.startswith('='):
                result=evaluate(expression(raw))
                if isinstance(result,list):raise FormulaError('#VALUE!','Use an aggregate function for a range.')
                if isinstance(result,(int,float)) and not isinstance(result,bool):result=numeric(result)
            elif raw in ('',None):result=None
            else:
                try:result=numeric(float(raw))
                except (ValueError,TypeError):result=raw
            cache[key]=result;return result
        except FormulaError as error:cache[key]=error;raise
        except ZeroDivisionError:raise FormulaError('#DIV/0!','Division by zero.')
        except (OverflowError,RecursionError,InvalidOperation):raise FormulaError('#NUM!','Numeric or calculation limit reached.')
        except (ValueError,TypeError,SyntaxError,TokenizerError):raise FormulaError('#VALUE!','Invalid formula or value.')
        finally:visiting.discard(key)
    def evaluate(node):
        if isinstance(node,ast.Constant) and isinstance(node.value,(str,int,float,bool)):return node.value
        if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.UAdd,ast.USub)):return numeric(evaluate(node.operand))*(1 if isinstance(node.op,ast.UAdd) else -1)
        if isinstance(node,ast.BinOp):
            left,right=evaluate(node.left),evaluate(node.right)
            if isinstance(node.op,ast.BitAnd):return str(left if left is not None else '')+str(right if right is not None else '')
            functions={ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv,ast.Pow:operator.pow}
            if type(node.op) not in functions:raise FormulaError('#NAME?','Unsupported operator.')
            left,right=numeric(left),numeric(right)
            if isinstance(node.op,ast.Pow) and abs(right)>100:raise FormulaError('#NUM!','Exponent exceeds the calculation limit.')
            return numeric(functions[type(node.op)](left,right))
        if isinstance(node,ast.Compare) and len(node.ops)==1:
            functions={ast.Eq:operator.eq,ast.NotEq:operator.ne,ast.Lt:operator.lt,ast.LtE:operator.le,ast.Gt:operator.gt,ast.GtE:operator.ge}
            if type(node.ops[0]) in functions:return functions[type(node.ops[0])](evaluate(node.left),evaluate(node.comparators[0]))
        if isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and not node.keywords:
            name=node.func.id.upper()
            if name=='ERROR' and len(node.args)==1 and isinstance(node.args[0],ast.Constant):raise FormulaError(str(node.args[0].value),'Referenced Excel error.')
            if name=='CELL' and len(node.args)==1 and isinstance(node.args[0],ast.Constant):return value(node.args[0].value)
            if name=='RANGE' and len(node.args)==2 and all(isinstance(n,ast.Constant) and isinstance(n.value,str) for n in node.args):
                r1,c1=coords(node.args[0].value);r2,c2=coords(node.args[1].value)
                if r2<r1 or c2<c1 or (r2-r1+1)*(c2-c1+1)>10000:raise FormulaError('#REF!','Range is reversed or exceeds 10,000 cells.')
                return [value(f'{get_column_letter(c+1)}{r+1}') for r in range(r1,r2+1) for c in range(c1,c2+1)]
            if name=='IF' and len(node.args) in (2,3):return evaluate(node.args[1]) if evaluate(node.args[0]) else evaluate(node.args[2]) if len(node.args)==3 else False
            if name=='IFERROR' and len(node.args)==2:
                try:return evaluate(node.args[0])
                except FormulaError:return evaluate(node.args[1])
                except (ValueError,TypeError,ArithmeticError):return evaluate(node.args[1])
            args=[]
            for arg in node.args:
                result=evaluate(arg);args.extend(result if isinstance(result,list) else [result])
            numbers=[numeric(v) for v in args if isinstance(v,(int,float)) and not isinstance(v,bool)]
            if name=='SUM':return sum(numbers)
            if name=='COUNT':return len(numbers)
            if name=='COUNTA':return sum(v not in ('',None) for v in args)
            if name in ('AVERAGE','MIN','MAX'):
                if not numbers:
                    if name=='AVERAGE':raise ZeroDivisionError()
                    return 0
                return {'AVERAGE':lambda:sum(numbers)/len(numbers),'MIN':lambda:min(numbers),'MAX':lambda:max(numbers)}[name]()
            if name=='ABS' and len(args)==1:return abs(numeric(args[0]))
            if name=='ROUND' and len(args)==2:
                digits=numeric(args[1])
                if digits!=int(digits) or abs(digits)>100:raise FormulaError('#NUM!','ROUND digits must be an integer between -100 and 100.')
                return float(Decimal(str(numeric(args[0]))).quantize(Decimal(1).scaleb(-int(digits)),rounding=ROUND_HALF_UP))
            if name=='AND':return all(bool(v) for v in args)
            if name=='OR':return any(bool(v) for v in args)
            if name=='NOT' and len(args)==1:return not bool(args[0])
        raise FormulaError('#NAME?','Unsupported function or expression.')
    return value


def calculate_cell(rows,address):
    result=evaluator(rows)(address)
    return 0 if result is None else result


def display(value,number_format=None):
    if value is None:return '0'
    if isinstance(value,bool):return 'TRUE' if value else 'FALSE'
    if isinstance(value,(int,float)):
        if not math.isfinite(value):return '#NUM!'
        if number_format=='0.00':return f'{value:.2f}'
        if number_format=='0.00%':return f'{value*100:.2f}%'
        if number_format=='#,##0.00':return f'{value:,.2f}'
        return f'{value:.12g}'
    return str(value)


class ResultDelegate(QStyledItemDelegate):
    def __init__(self,editor):super().__init__(editor.grid);self.editor=editor
    def initStyleOption(self,option,index):
        super().initStyleOption(option,index);raw=index.data(Qt.EditRole)
        if isinstance(raw,str) and raw.startswith('=') and not self.editor.show_formulas:
            result=self.editor.formula_results.get((index.row(),index.column()),'…');option.text=result.code if isinstance(result,FormulaError) else display(result,index.data(NUMBER_FORMAT_ROLE))
        elif index.data(NUMBER_FORMAT_ROLE) in ('0.00','0.00%','#,##0.00'):
            try:option.text=display(float(raw),index.data(NUMBER_FORMAT_ROLE))
            except (ValueError,TypeError):pass


def recalculate(editor):
    editor.recalc_timer.stop();rows=editor._rows();compute=evaluator(rows);results={}
    with QSignalBlocker(editor.grid):
        for r,row in enumerate(rows):
            for c,raw in enumerate(row):
                if not raw.startswith('='):
                    if (r,c) in editor.formula_results:
                        item=editor.grid.item(r,c)
                        if item:item.setToolTip('');item.setData(Qt.AccessibleTextRole,None)
                    continue
                try:result=compute(f'{get_column_letter(c+1)}{r+1}')
                except FormulaError as error:result=error
                results[r,c]=result;item=editor.grid.item(r,c)
                item.setToolTip(raw+'\n'+(str(result) if isinstance(result,FormulaError) else display(result)))
                item.setData(Qt.AccessibleTextRole,raw if editor.show_formulas else result.code if isinstance(result,FormulaError) else display(result))
    editor.formula_results=results;editor.grid.viewport().update()


def computed_value(editor,row,column):
    if editor.recalc_timer.isActive():recalculate(editor)
    item=editor.grid.item(row,column);raw=item.text() if item else ''
    if raw.startswith('='):
        result=editor.formula_results.get((row,column))
        if isinstance(result,FormulaError):raise result
        return 0 if result is None else result
    try:return float(raw) if raw else 0
    except ValueError:return raw


def setup_formulas(editor):
    editor.show_formulas=False;editor.formula_results={};editor.recalc_timer=QTimer(editor);editor.recalc_timer.setSingleShot(True);editor.recalc_timer.setInterval(100)
    editor.recalc_timer.timeout.connect(lambda:recalculate(editor));editor.grid.itemChanged.connect(lambda _:editor.recalc_timer.start())
    editor.grid.setItemDelegate(ResultDelegate(editor))
    command(editor,'Formulas','Recalculate',lambda:recalculate(editor),'calculate','F9')
    action=command(editor,'Formulas','Show formulas',lambda:None,'calculate');action.setCheckable(True)
    def toggle(checked):editor.show_formulas=checked;recalculate(editor)
    action.toggled.connect(toggle)
    def autosum():
        ranges=editor.grid.selectedRanges()
        if len(ranges)!=1 or ranges[0].columnCount()!=1:raise ValueError('Select one column of cells to sum.')
        area=ranges[0];r,c=area.bottomRow()+1,area.leftColumn()
        if r>=1048576:raise ValueError('No row is available below the selection.')
        if r<editor.grid.rowCount() and editor.grid.item(r,c) and editor.grid.item(r,c).text():raise ValueError('The cell below the selection is not empty; choose another range.')
        editor.grid.setRowCount(max(editor.grid.rowCount(),r+1));letter=get_column_letter(c+1);editor.grid.setItem(r,c,QTableWidgetItem(f'=SUM({letter}{area.topRow()+1}:{letter}{area.bottomRow()+1})'));editor.grid.setCurrentCell(r,c)
    command(editor,'Formulas','AutoSum',autosum,'calculate','Alt+=')
