import unittest
import tempfile
from pathlib import Path
from windows_integration import registration_values, MEDIA_EXTENSIONS, WORKSPACE_EXTENSIONS, launcher_name


class WindowsIntegrationTest(unittest.TestCase):
    def test_quoted_media_handler_does_not_set_defaults(self):
        executable = Path('build/App with spaces/LocalFileConverter.exe').resolve()
        values = registration_values(executable)
        commands = [v for path, name, v in values if path.endswith('open\\command')]
        self.assertEqual(commands[:2], [f'"{executable}" --play "%1"'] * 2)
        for workspace, extensions in WORKSPACE_EXTENSIONS.items():
            command = f'"{executable.parent / launcher_name(workspace)}" "%1"'
            self.assertEqual(commands.count(command), 2)
            for extension in extensions:
                self.assertIn(('Software\\Classes\\'+extension+'\\OpenWithProgids', 'DrunkPeanuts.'+workspace, ''), values)
        for extension in MEDIA_EXTENSIONS:
            self.assertIn(('Software\\Classes\\' + extension + '\\OpenWithProgids', 'FreeConvert.Media', ''), values)
        self.assertFalse(any('UserChoice' in path or name == 'ProgId' for path, name, _ in values))
        self.assertIn(('Software\\Classes\\FreeConvert.Media', 'NoOpenWith', ''), values)
        self.assertIn(('Software\\Classes\\Applications\\LocalFileConverter.exe', 'NoOpenWith', ''), values)

    def test_open_with_dispatches_each_office_file_without_overwriting_open_work(self):
        import os
        os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
        from PySide6.QtWidgets import QApplication
        from unittest.mock import MagicMock, patch
        from app import Window
        application = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as folder:
            first, second = Path(folder)/'First file.docx', Path(folder)/'Second file.docx'
            first.touch();second.touch()
            window = Window()
            with patch('office_suite.WORKSPACES', {'Pages':MagicMock(side_effect=[MagicMock(), MagicMock()])}), patch.object(window, '_workspace_templates'):
                window.open_arguments(['--open-workspace', 'Pages', str(first), str(second)])
                opened = list(window.workspaces.values())
                self.assertEqual(len(opened), 2)
                opened[0].open.assert_called_once_with(first.resolve())
                opened[1].open.assert_called_once_with(second.resolve())
            window.workspaces.clear();window.close()

    def test_templates_are_editable_copies_and_never_overwrite(self):
        from office_templates import ASSETS, TEMPLATES, create_document
        from docx import Document
        from openpyxl import load_workbook
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as folder:
            for workspace, names in TEMPLATES.items():
                for name in names:
                    target = create_document(workspace, name, Path(folder)/name)
                    self.assertEqual(target.read_bytes(), (ASSETS/'templates'/name).read_bytes())
                    with self.assertRaises(FileExistsError):
                        create_document(workspace, name, target)
            self.assertIn('Report title', [p.text for p in Document(Path(folder)/'Report.docx').paragraphs])
            book = load_workbook(Path(folder)/'Budget.xlsx')
            self.assertEqual(book.active['D2'].value, '=B2-C2')
            book.close()
            self.assertEqual(len(Presentation(Path(folder)/'Presentation.pptx').slides), 4)
            import sqlite3
            from contextlib import closing
            with closing(sqlite3.connect(Path(folder)/'Contacts.sqlite')) as db:
                self.assertEqual(db.execute("SELECT value FROM _freeconvert_meta WHERE key='owner'").fetchone(), ('Organize',))


if __name__ == '__main__':
    unittest.main()
