"""Core editing commands built on the existing native document models."""
import copy
from PySide6.QtCore import Qt
from PySide6.QtGui import QAction,QFont,QTextCharFormat,QTextCursor,QTextDocument,QTextListFormat,QTextFormat,QColor
from PySide6.QtWidgets import (QDialog,QVBoxLayout,QFormLayout,QLineEdit,QCheckBox,QPushButton,
    QLabel,QInputDialog,QColorDialog,QDialogButtonBox,QDoubleSpinBox)
from ui import icon,palette


def command(window,menu,label,callback,glyph='settings',shortcut=None):
    menus={a.text().replace('&',''):a.menu() for a in window.menuBar().actions()}
    target=menus.get(menu) or window.menuBar().addMenu(menu)
    action=QAction(icon(glyph,palette()['accent']),label,window);action.setProperty('iconName',glyph)
    def run():
        if hasattr(window,'report'):window.report(callback)
        else:callback()
    action.triggered.connect(run)
    if shortcut:action.setShortcut(shortcut)
    target.addAction(action);return action


def replace_text(editor,query,replacement,case=False,whole=False):
    if not query:raise ValueError('Enter text to find.')
    flags=QTextDocument.FindFlags()
    if case:flags|=QTextDocument.FindCaseSensitively
    if whole:flags|=QTextDocument.FindWholeWords
    cursor=QTextCursor(editor.document());cursor.beginEditBlock();count=0
    while True:
        match=editor.document().find(query,cursor,flags)
        if match.isNull():break
        match.insertText(replacement);cursor.setPosition(match.position());count+=1
    cursor.endEditBlock();return count


def find_dialog(window,editor):
    dialog=QDialog(window);dialog.setAttribute(Qt.WA_DeleteOnClose);dialog.setWindowTitle('Find and replace');layout=QVBoxLayout(dialog);form=QFormLayout();query=QLineEdit();replacement=QLineEdit();form.addRow('Find',query);form.addRow('Replace with',replacement);layout.addLayout(form)
    case=QCheckBox('Match case');whole=QCheckBox('Whole words');layout.addWidget(case);layout.addWidget(whole);status=QLabel();layout.addWidget(status)
    def find():
        if not query.text():return
        flags=QTextDocument.FindFlags()
        if case.isChecked():flags|=QTextDocument.FindCaseSensitively
        if whole.isChecked():flags|=QTextDocument.FindWholeWords
        if not editor.find(query.text(),flags):
            cursor=editor.textCursor();cursor.movePosition(QTextCursor.Start);editor.setTextCursor(cursor)
            if not editor.find(query.text(),flags):status.setText('No match');return
        status.setText('Match selected')
    def replace_all():
        if query.text():status.setText(f'{replace_text(editor,query.text(),replacement.text(),case.isChecked(),whole.isChecked())} replacements')
    for label,callback in [('Find next',find),('Replace all',replace_all)]:
        button=QPushButton(label);button.clicked.connect(callback);layout.addWidget(button)
    dialog.resize(390,230);dialog.show();query.setFocus();return dialog


def heading(pages,level):
    cursor=pages.editor.textCursor();block=cursor.blockFormat();block.setHeadingLevel(level);block.setTopMargin(12 if level else 0);block.setBottomMargin(6);cursor.setBlockFormat(block)
    fmt=QTextCharFormat();fmt.setFontWeight(QFont.Bold if level else QFont.Normal);fmt.setFontPointSize({0:12,1:26,2:20,3:16}[level]);cursor.select(QTextCursor.BlockUnderCursor);cursor.mergeCharFormat(fmt);pages.rich_changed=True


def paragraph_spacing(pages,value=None):
    if value is None:
        value,ok=QInputDialog.getDouble(pages,'Line spacing','Multiple',1.15,.8,3,2)
        if not ok:return
    cursor=pages.editor.textCursor();fmt=cursor.blockFormat();fmt.setLineHeight(value*100,1);cursor.setBlockFormat(fmt);pages.rich_changed=True


def page_setup(pages,settings=None):
    if settings is None:
        dialog=QDialog(pages);dialog.setWindowTitle('Page setup');layout=QFormLayout(dialog)
        from PySide6.QtWidgets import QComboBox
        paper=QComboBox();paper.addItems(['A4','Letter']);orientation=QComboBox();orientation.addItems(['Portrait','Landscape']);margin=QDoubleSpinBox();margin.setRange(5,60);margin.setSuffix(' mm')
        current=pages.page_settings or dict(paper='A4',landscape=False,margin=20)
        paper.setCurrentText(current['paper']);orientation.setCurrentIndex(int(current['landscape']));margin.setValue(current['margin']);layout.addRow('Paper',paper);layout.addRow('Orientation',orientation);layout.addRow('Margins',margin)
        buttons=QDialogButtonBox(QDialogButtonBox.Ok|QDialogButtonBox.Cancel);buttons.accepted.connect(dialog.accept);buttons.rejected.connect(dialog.reject);layout.addRow(buttons)
        if dialog.exec()!=QDialog.Accepted:return
        settings=dict(paper=paper.currentText(),landscape=bool(orientation.currentIndex()),margin=margin.value())
    if settings['paper'] not in ('A4','Letter') or not 5<=settings['margin']<=60:raise ValueError('Invalid page settings.')
    pages.page_settings=dict(settings);pages.rich_changed=True;pages.editor.document().setModified(True)


def setup_pages_tools(pages):
    pages.page_settings=None
    for label,level in [('Normal',0),('Heading 1',1),('Heading 2',2),('Heading 3',3)]:command(pages,'Home',label,lambda n=level:heading(pages,n),'title')
    command(pages,'Home','Numbering',lambda:pages.editor.textCursor().createList(QTextListFormat.ListDecimal),'bullets')
    command(pages,'Home','Line spacing',lambda:paragraph_spacing(pages),'settings')
    def char_change(kind):
        fmt=QTextCharFormat();current=pages.editor.currentCharFormat()
        if kind=='Strike':fmt.setFontStrikeOut(not current.fontStrikeOut())
        elif kind=='Super':fmt.setVerticalAlignment(QTextCharFormat.AlignNormal if current.verticalAlignment()==QTextCharFormat.AlignSuperScript else QTextCharFormat.AlignSuperScript)
        elif kind=='Sub':fmt.setVerticalAlignment(QTextCharFormat.AlignNormal if current.verticalAlignment()==QTextCharFormat.AlignSubScript else QTextCharFormat.AlignSubScript)
        elif kind=='Highlight':
            color=QColorDialog.getColor(QColor('#ffdd22'),pages)
            if not color.isValid():return
            fmt.setBackground(color)
        else:
            cursor=pages.editor.textCursor();cursor.setCharFormat(QTextCharFormat());pages.rich_changed=True;return
        pages.editor.mergeCurrentCharFormat(fmt);pages.rich_changed=True
    for label,kind,glyph in [('Strikethrough','Strike','strike'),('Superscript','Super','superscript'),('Subscript','Sub','subscript'),('Highlight','Highlight','highlight'),('Clear formatting','Clear','eraser')]:command(pages,'Home',label,lambda k=kind:char_change(k),glyph)
    command(pages,'Review','Find and replace',lambda:find_dialog(pages,pages.editor),'find','Ctrl+H')
    command(pages,'Layout','Page setup',lambda:page_setup(pages),'settings')
    command(pages,'View','Zoom in',lambda:pages.editor.zoomIn(1),'add','Ctrl+=')
    command(pages,'View','Zoom out',lambda:pages.editor.zoomOut(1),'settings','Ctrl+-')


def duplicate_slide(editor):
    def duplicate():
        from pathlib import PurePosixPath
        from pptx.opc.packuri import PackURI
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        package=editor.deck.part.package;used={str(p.partname) for p in package.iter_parts()};cloned={}
        def clone_part(part):
            if part in cloned:return cloned[part]
            name=PurePosixPath(str(part.partname));index=1
            while str(name.with_name(name.stem+f'_copy{index}'+name.suffix)) in used:index+=1
            target=str(name.with_name(name.stem+f'_copy{index}'+name.suffix));used.add(target)
            result=type(part).load(PackURI(target),part.content_type,package,part.blob);cloned[part]=result
            mapping={}
            for rel in part.rels.values():mapping[rel.rId]=result.relate_to(rel.target_ref if rel.is_external else clone_part(rel.target_part),rel.reltype,rel.is_external)
            if hasattr(result,'_element'):
                for node in result._element.iter():
                    for key,value in list(node.attrib.items()):
                        if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') and value in mapping:node.set(key,mapping[value])
            return result
        source=editor.deck.slides[editor.index];slide=editor.deck.slides.add_slide(editor.deck.slide_layouts[6])
        def copy_node(source_node):
            element=copy.deepcopy(source_node)
            for node in element.iter():
                for key,value in list(node.attrib.items()):
                    if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                        rel=source.part.rels[value]
                        target=rel.target_ref if rel.is_external else clone_part(rel.target_part) if rel.reltype==RT.CHART else rel.target_part
                        node.set(key,slide.part.relate_to(target,rel.reltype,rel.is_external))
            return element
        for shape in source.shapes:
            element=copy_node(shape.element)
            slide.shapes._spTree.insert_element_before(element,'p:extLst')
            target_shape=slide.shapes[-1]
            target_shape.left=shape.left;target_shape.top=shape.top;target_shape.width=shape.width;target_shape.height=shape.height
        for tag in ('p:bg','p:transition','p:timing','p:extLst'):
            node=source._element.cSld.find(tag,source._element.nsmap) if tag=='p:bg' else source._element.find(tag,source._element.nsmap)
            if node is not None:
                if tag=='p:bg':slide._element.cSld.insert(0,copy_node(node))
                else:slide._element.append(copy_node(node))
        if source.has_notes_slide:
            # Copy the native notes XML, remapping its relationships, so notes stay independent.
            notes=source.notes_slide
            target=slide.notes_slide
            element=copy.deepcopy(notes._element)
            for node in element.iter():
                for key,value in list(node.attrib.items()):
                    if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                        rel=notes.part.rels[value]
                        node.set(key,target.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,rel.reltype,rel.is_external))
            target._element[:]=list(element)
        editor.index=len(editor.deck.slides)-1
    editor.change(duplicate);editor.refresh()


def align_object(editor,edge):
    shape=editor.selected_shape()
    if shape is None:raise ValueError('Select a slide object first.')
    width,height=editor.deck.slide_width,editor.deck.slide_height
    def align():
        if edge=='Left':shape.left=0
        elif edge=='Center':shape.left=max(0,(width-shape.width)//2)
        elif edge=='Right':shape.left=max(0,width-shape.width)
        elif edge=='Top':shape.top=0
        elif edge=='Middle':shape.top=max(0,(height-shape.height)//2)
        else:shape.top=max(0,height-shape.height)
    editor.change(align)


def setup_slide_tools(editor):
    command(editor,'Slide','Duplicate slide',lambda:duplicate_slide(editor),'copy','Ctrl+Shift+D')
    for edge in ('Left','Center','Right','Top','Middle','Bottom'):command(editor,'Arrange','Align '+edge.lower(),lambda e=edge:align_object(editor,e),'center')
    command(editor,'Animations','Fade on click',lambda:editor.report(lambda:editor.animate_object(editor.selected_shape(),.6)),'timer')
    command(editor,'Animations','Remove animation',lambda:editor.report(lambda:editor.animate_object(editor.selected_shape(),None)),'delete')
    command(editor,'View','Fit slide',lambda:editor.slide_page.fitInView(editor.scene.sceneRect(),Qt.KeepAspectRatio),'fullscreen')
