"""Editable, single-column PDF reconstruction using the bundled PDFium engine."""
import ctypes,io,re
from pathlib import Path
import pypdfium2 as pdfium
from pypdfium2 import raw
from docx import Document
from docx.shared import Pt
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _font(text,index):
    flags=ctypes.c_int();length=raw.FPDFText_GetFontInfo(text,index,None,0,ctypes.byref(flags));name='Arial'
    if 0<length<1024:
        buffer=ctypes.create_string_buffer(length);raw.FPDFText_GetFontInfo(text,index,buffer,length,ctypes.byref(flags));name=buffer.value.decode('utf-8',errors='replace')
    name=re.sub(r'^[A-Z]{6}\+','',name)
    family='Times New Roman' if 'Times' in name else 'Arial' if 'Arial' in name or 'Helvetica' in name else 'Calibri'
    size=raw.FPDFText_GetFontSize(text,index)
    return family,max(4,min(96,size or 11)),'bold' in name.lower(),'italic' in name.lower() or 'oblique' in name.lower()


def _lines(text,height):
    result=[];chars=[];boxes=[];fonts=[]
    def finish():
        value=''.join(chars).strip()
        if value and boxes:
            left=min(b[0] for b in boxes);right=max(b[2] for b in boxes);top=height-max(b[3] for b in boxes)
            font=max(fonts,key=lambda f:f[1]);result.append((top,left,right,value,font))
        chars.clear();boxes.clear();fonts.clear()
    for index in range(text.count_chars()):
        char=text.get_text_range(index,1)
        if char in ('\r','\n'):finish();continue
        chars.append(char)
        if char.strip():
            box=text.get_charbox(index)
            if box[2]>box[0] and box[3]>box[1]:boxes.append(box);fonts.append(_font(text,index))
    finish();return sorted(result,key=lambda line:(round(line[0],1),line[1]))


def _anchor(inline,x,y,identifier):
    """Keep a native editable picture at its PDF position without moving body text."""
    anchor=OxmlElement('wp:anchor')
    for key,value in dict(distT='0',distB='0',distL='0',distR='0',simplePos='0',relativeHeight=str(identifier),behindDoc='1',locked='0',layoutInCell='1',allowOverlap='1').items():anchor.set(key,value)
    simple=OxmlElement('wp:simplePos');simple.set('x','0');simple.set('y','0');anchor.append(simple)
    for name,relative,offset in [('positionH','page',x),('positionV','page',y)]:
        node=OxmlElement('wp:'+name);node.set('relativeFrom',relative);value=OxmlElement('wp:posOffset');value.text=str(round(offset*12700));node.append(value);anchor.append(node)
    anchor.append(inline.find(qn('wp:extent')));anchor.append(OxmlElement('wp:wrapNone'))
    for name in ('wp:docPr','wp:cNvGraphicFramePr','a:graphic'):
        node=inline.find(qn(name))
        if node is not None:anchor.append(node)
    inline.getparent().replace(inline,anchor)


def editable_docx(source,target):
    document=Document();text_count=0
    with pdfium.PdfDocument(str(source)) as pdf:
        for index in range(len(pdf)):
            page=pdf[index]
            try:
                if page.get_rotation()!=0:raise ValueError('Editable reconstruction currently requires unrotated PDF pages. Rotate the PDF first or choose Preserve pages.')
                width,height=page.get_size();section=document.sections[0] if index==0 else document.add_section(WD_SECTION_START.NEW_PAGE)
                section.page_width=Pt(width);section.page_height=Pt(height)
                section.left_margin=section.right_margin=section.top_margin=section.bottom_margin=Pt(0)
                images=document.add_paragraph();images.paragraph_format.space_before=images.paragraph_format.space_after=Pt(0);images.paragraph_format.line_spacing=Pt(1)
                for obj in page.get_objects():
                    if obj.type!=raw.FPDF_PAGEOBJ_IMAGE:continue
                    left,bottom,right,top=obj.get_bounds();w,h=right-left,top-bottom
                    if w<=0 or h<=0:continue
                    bitmap=obj.get_bitmap(render=True)
                    try:
                        blob=io.BytesIO()
                        with bitmap.to_pil() as image:image.save(blob,format='PNG')
                        blob.seek(0)
                        picture=images.add_run().add_picture(blob,width=Pt(w),height=Pt(h));_anchor(picture._inline,max(0,left),max(0,height-top),picture._inline.docPr.id)
                    finally:
                        # The PIL view is closed and PNG bytes own their storage before freeing PDFium's buffer.
                        bitmap.warn_on_close=False
                        bitmap.close()
                text=page.get_textpage()
                try:lines=_lines(text,height)
                finally:text.close()
                if not lines:raise ValueError('This PDF page has no selectable text. Use OCR first, or choose Preserve pages (not editable).')
                # ponytail: single-column line reconstruction; complex columns/tables need a layout parser.
                cursor=1.0
                for top,left,right,value,font in lines:
                    paragraph=document.add_paragraph();fmt=paragraph.paragraph_format
                    fmt.left_indent=Pt(max(0,left));fmt.right_indent=Pt(max(0,width-right-8));fmt.space_before=Pt(max(0,top-cursor));fmt.space_after=Pt(0);fmt.line_spacing=Pt(font[1]*1.05)
                    run=paragraph.add_run(value);run.font.name=font[0];run.font.size=Pt(font[1]);run.bold=font[2];run.italic=font[3]
                    cursor=max(cursor,top)+font[1]*1.05;text_count+=len(value)
            finally:page.close()
    if not text_count:raise ValueError('No editable text was found in the PDF.')
    created=False
    try:
        with Path(target).open('xb') as output:
            created=True
            document.save(output)
    except Exception:
        if created:Path(target).unlink(missing_ok=True)
        raise
