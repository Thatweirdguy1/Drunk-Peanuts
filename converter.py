"""Local file conversion. External tools run without a shell and never replace inputs."""

from __future__ import annotations

import csv
import io
import math
import os
import re
import shutil
import stat
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile
from pathlib import Path

import imageio_ffmpeg
from native_office import DOCUMENTS, PRESENTATIONS, SPREADSHEETS, advanced_content, convert_office, csv_delimiter, targets
from office_pdf import (LEGACY_FORMATS, LIBREOFFICE_FORMATS, MODERN_FOR_LEGACY,
                        available_engines, convert_legacy, export_pdf, resolve_engine)
import pillow_heif
import pypdfium2 as pdfium
import rawpy
import rapidocr
from rapidocr import ModelType, OCRVersion, RapidOCR
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageOps, ImageSequence, UnidentifiedImageError
from PySide6.QtCore import QBuffer, QByteArray, QIODevice
from PySide6.QtGui import QImage, QPainter
from PySide6.QtSvg import QSvgRenderer
from pypdf import PdfReader, PdfWriter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

pillow_heif.register_heif_opener()
Image.MAX_IMAGE_PIXELS = 100_000_000

IMAGE = {"jpg", "jpeg", "png", "heic", "heif", "tif", "tiff", "gif", "bmp", "webp", "svg", "cr2", "nef", "arw", "ico"}
RASTER = {"jpg", "png", "webp", "tiff", "bmp", "gif", "ico", "pdf"}
VIDEO = {"mp4", "mov", "avi", "mpeg", "mpg", "mkv", "webm"}
AUDIO = {"mp3", "wav", "aiff", "m4a", "aac", "flac", "ogg"}
OFFICE = ((DOCUMENTS | PRESENTATIONS | SPREADSHEETS) - {"txt", "csv"}) | LEGACY_FORMATS
FFMPEG_VIDEO = {"mp4", "mov", "avi", "mpeg", "mkv", "webm", "gif", "mp3", "wav", "m4a", "aac", "flac", "ogg"}
FFMPEG_AUDIO = {"mp3", "wav", "aiff", "m4a", "aac", "flac", "ogg"}


def kind(path: Path) -> str:
    suffix = path.suffix.lower().lstrip(".")
    if suffix in IMAGE:
        return "Image"
    if suffix == "pdf":
        return "PDF"
    if suffix in VIDEO:
        return "Video"
    if suffix in AUDIO:
        return "Audio"
    if suffix in OFFICE:
        return "Document"
    if suffix == "csv":
        return "Spreadsheet"
    if suffix == "txt":
        return "Text"
    return "Unsupported"


def available_formats(path: Path) -> list[str]:
    ext = path.suffix.lower().lstrip(".")
    group = kind(path)
    if group == "Image":
        choices = RASTER | {"svg"} if ext == "svg" else RASTER
        if ext in {"gif", "webp", "tif", "tiff"} and path.is_file():
            with Image.open(path) as image:
                if getattr(image, "n_frames", 1) > 1:
                    choices &= {"gif", "webp", "tiff", "pdf"}
    elif group == "PDF":
        choices = {"png", "jpg", "webp", "txt", "docx"}
    elif group == "Video":
        choices = FFMPEG_VIDEO
    elif group == "Audio":
        choices = FFMPEG_AUDIO
    elif ext in LEGACY_FORMATS:
        choices = {"pdf", MODERN_FOR_LEGACY[ext]} if available_engines(path) else set()
    elif ext in DOCUMENTS | SPREADSHEETS | PRESENTATIONS:
        choices = targets(ext) | ({"txt"} if ext == "csv" else set())
    else:
        choices = set()
    return sorted(choices - {ext})


def ffmpeg_binary() -> str:
    return imageio_ffmpeg.get_ffmpeg_exe()


def _run(command: list[str], timeout: int = 900) -> None:
    result = subprocess.run(command, capture_output=True, text=True, timeout=timeout, shell=False)
    if result.returncode:
        raise RuntimeError((result.stderr or result.stdout)[-700:].strip() or "Converter failed")


def _image(path: Path) -> Image.Image:
    ext = path.suffix.lower().lstrip(".")
    if ext == "svg":
        data = path.read_bytes()
        root = ET.fromstring(data)
        if any(value and not value.startswith("#") for node in root.iter() for key, value in node.attrib.items() if key.endswith("href")):
            raise ValueError("SVG files with external resources are not supported.")
        renderer = QSvgRenderer(QByteArray(data))
        if not renderer.isValid():
            raise ValueError("Invalid SVG file.")
        size = renderer.defaultSize()
        picture = QImage(max(1, size.width()), max(1, size.height()), QImage.Format_ARGB32)
        picture.fill(0)
        painter = QPainter(picture)
        renderer.render(painter)
        painter.end()
        buffer = QBuffer()
        buffer.open(QIODevice.WriteOnly)
        picture.save(buffer, "PNG")
        return Image.open(io.BytesIO(bytes(buffer.data())))
    if ext in {"cr2", "nef", "arw"}:
        with rawpy.imread(str(path)) as raw:
            return Image.fromarray(raw.postprocess(use_camera_wb=True))
    return Image.open(path)


def _save_image(source: Path, target: Path, fmt: str) -> None:
    with _image(source) as opened:
        if getattr(opened, "n_frames", 1) > 1:
            if fmt not in {"pdf", "gif", "webp", "tiff"}:
                raise ValueError("Choose GIF, WebP, TIFF, or PDF to keep every frame.")
            frames, durations = [], []
            for frame in ImageSequence.Iterator(opened):
                durations.append(frame.info.get("duration", opened.info.get("duration", 100)))
                frames.append(ImageOps.exif_transpose(frame.copy()).convert("RGB" if fmt in {"pdf", "tiff"} else "RGBA"))
            options = {"save_all": True, "append_images": frames[1:]}
            if fmt in {"gif", "webp"}:
                options.update(duration=durations, loop=opened.info.get("loop", 0))
            frames[0].save(target, {"tiff": "TIFF"}.get(fmt, fmt.upper()), **options)
        else:
            frame = ImageOps.exif_transpose(opened)
            if fmt in {"jpg", "bmp"}:
                background = Image.new("RGB", frame.size, "white")
                if frame.mode == "RGBA":
                    background.paste(frame, mask=frame.getchannel("A"))
                else:
                    background.paste(frame.convert("RGB"))
                frame = background
            frame.save(target, {"jpg": "JPEG", "tiff": "TIFF"}.get(fmt, fmt.upper()))


def _pdf_to_image(source: Path, target: Path, fmt: str) -> None:
    pdf = pdfium.PdfDocument(str(source))
    try:
        if any(target.with_name(f"{target.stem}-{index + 1}{target.suffix}").exists() for index in range(1, len(pdf))):
            raise FileExistsError("A numbered PDF page output already exists.")
        # Multi-page image exports use numbered files alongside the selected first-page name.
        for index in range(len(pdf)):
            page = pdf[index]
            image = page.render(scale=2).to_pil()
            output = target if index == 0 else target.with_name(f"{target.stem}-{index + 1}{target.suffix}")
            if fmt == "jpg":
                image = image.convert("RGB")
            image.save(output, {"jpg": "JPEG"}.get(fmt, fmt.upper()))
            page.close()
    finally:
        pdf.close()


def _csv_to_xlsx(source: Path, target: Path) -> None:
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Sheet1")
    with source.open("r", encoding="utf-8-sig", newline="") as stream:
        for row in csv.reader(stream, delimiter=csv_delimiter(source)):
            sheet.append(row)
    book.save(target)


def _xlsx_to_csv(source: Path, target: Path) -> None:
    book = load_workbook(source, read_only=True, data_only=True)
    try:
        with target.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            for row in book.active.iter_rows(values_only=True):
                writer.writerow(["" if value is None else value for value in row])
    finally:
        book.close()


def _text_to_pdf(source: Path, target: Path) -> None:
    document = canvas.Canvas(str(target))
    text = document.beginText(48, 800)
    for line in source.read_text(encoding="utf-8-sig").splitlines():
        if text.getY() < 48:
            document.drawText(text)
            document.showPage()
            text = document.beginText(48, 800)
        text.textLine(line[:120])
    document.drawText(text)
    document.save()


def _text_to_docx(source: Path, target: Path) -> None:
    document = Document()
    for line in source.read_text(encoding="utf-8-sig").splitlines():
        document.add_paragraph(line)
    document.save(target)


def _pdf_text(source: Path) -> str:
    return "\n\n".join(page.extract_text() or "" for page in PdfReader(str(source)).pages)


def _pdf_to_docx_visual(source: Path, target: Path) -> None:
    from docx.enum.section import WD_SECTION_START
    from docx.shared import Pt

    document = Document()
    pdf = pdfium.PdfDocument(str(source))
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            try:
                width, height = page.get_size()
                section = document.sections[0] if index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
                section.page_width, section.page_height = Pt(width), Pt(height)
                section.left_margin = section.right_margin = Pt(0)
                section.top_margin = section.bottom_margin = Pt(0)
                picture = io.BytesIO()
                page.render(scale=1.5).to_pil().save(picture, format="PNG")
                picture.seek(0)
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
                paragraph.add_run().add_picture(picture, width=Pt(max(1, width - 6)), height=Pt(max(1, height - 6)))
            finally:
                page.close()
    finally:
        pdf.close()
    document.save(target)


def pdf_merge(sources: list[Path], target: Path) -> None:
    if len(sources) < 2 or any(path.suffix.lower() != ".pdf" for path in sources):
        raise ValueError("Select at least two PDFs to merge.")
    if target.exists() or target in sources:
        raise FileExistsError(f"Output already exists: {target.name}")
    writer = PdfWriter()
    created = False
    try:
        for source in sources:
            writer.append(str(source))
        with target.open("xb") as output:
            created = True
            writer.write(output)
    except Exception:
        if created:
            target.unlink(missing_ok=True)
        raise
    finally:
        writer.close()


def pdf_split(source: Path, folder: Path) -> None:
    if source.suffix.lower() != ".pdf" or folder.exists():
        raise ValueError("Choose a PDF and a new output folder.")
    pages = PdfReader(str(source)).pages
    folder.mkdir(parents=True, exist_ok=False)
    try:
        for index, page in enumerate(pages, 1):
            writer = PdfWriter()
            writer.add_page(page)
            with (folder / f"page-{index}.pdf").open("xb") as output:
                writer.write(output)
            writer.close()
    except Exception:
        shutil.rmtree(folder)
        raise


def pdf_reorder(source: Path, target: Path, order: list[int]) -> None:
    pages = PdfReader(str(source)).pages
    if sorted(order) != list(range(1, len(pages) + 1)):
        raise ValueError(f"Enter each page number once, from 1 to {len(pages)}.")
    if target.exists() or target == source:
        raise FileExistsError(f"Output already exists: {target.name}")
    writer = PdfWriter()
    created = False
    try:
        for number in order:
            writer.add_page(pages[number - 1])
        with target.open("xb") as output:
            created = True
            writer.write(output)
    except Exception:
        if created:
            target.unlink(missing_ok=True)
        raise
    finally:
        writer.close()


def pdf_rotate(source: Path, target: Path, degrees: int) -> None:
    if degrees not in {90, 180, 270}:
        raise ValueError("Choose 90, 180, or 270 degrees.")
    if target.exists() or target == source:
        raise FileExistsError(f"Output already exists: {target.name}")
    writer = PdfWriter()
    created = False
    try:
        for page in PdfReader(str(source)).pages:
            writer.add_page(page.rotate(degrees))
        with target.open("xb") as output:
            created = True
            writer.write(output)
    except Exception:
        if created:
            target.unlink(missing_ok=True)
        raise
    finally:
        writer.close()


def edit_image(source: Path, target: Path, action: str, values: tuple[int, ...] = ()) -> None:
    if kind(source) != "Image" or target.exists() or source == target:
        raise ValueError("Choose an image and a new output path.")
    with _image(source) as opened:
        if getattr(opened, "n_frames", 1) > 1:
            raise ValueError("Edit animated images frame by frame after conversion.")
        image = ImageOps.exif_transpose(opened).copy()
    if action == "resize":
        width, height = values
        if width < 1 or height < 1:
            raise ValueError("Width and height must be positive.")
        image.thumbnail((width, height), Image.Resampling.LANCZOS)
    elif action == "crop":
        left, top, width, height = values
        if min(left, top) < 0 or min(width, height) < 1 or left + width > image.width or top + height > image.height:
            raise ValueError("Crop rectangle must fit inside the image.")
        image = image.crop((left, top, left + width, top + height))
    elif action == "rotate":
        degrees, = values
        if degrees not in {90, 180, 270}:
            raise ValueError("Choose 90, 180, or 270 degrees.")
        image = image.rotate(-degrees, expand=True)
    elif action == "strip_metadata":
        pass
    else:
        raise ValueError("Unknown image tool.")
    image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
    if target.suffix.lower() in {".jpg", ".jpeg", ".bmp"} and "A" in image.getbands():
        background = Image.new("RGB", image.size, "white")
        background.paste(image, mask=image.getchannel("A"))
        image = background
    # Pixel-only copy prevents EXIF, GPS, comments, and other source metadata from leaking.
    image = Image.frombytes(image.mode, image.size, image.tobytes())
    created = False
    try:
        with target.open("xb") as output:
            created = True
            image.save(output, format={".jpg": "JPEG", ".jpeg": "JPEG", ".tif": "TIFF", ".tiff": "TIFF"}.get(target.suffix.lower(), target.suffix[1:].upper()))
    except Exception:
        if created:
            target.unlink(missing_ok=True)
        raise


def convert(source: Path, target: Path, engine: str = "built-in") -> None:
    source, target = Path(source).resolve(strict=True), Path(target).resolve()
    if not source.is_file() or source == target:
        raise ValueError("Choose a file and a different output path.")
    fmt = target.suffix.lower().lstrip(".")
    if fmt not in available_formats(source):
        raise ValueError(f"{source.suffix} to .{fmt} is not supported.")
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        raise FileExistsError(f"Output already exists: {target.name}")
    office_pdf = fmt == "pdf" and source.suffix.lower().lstrip(".") in LIBREOFFICE_FORMATS
    legacy_modern = source.suffix.lower().lstrip(".") in LEGACY_FORMATS and fmt == MODERN_FOR_LEGACY[source.suffix.lower().lstrip(".")]
    selected_engine = resolve_engine(source, engine) if office_pdf else "built-in"
    pdf_visual = source.suffix.lower() == ".pdf" and fmt == "docx" and engine == "visual"
    if engine not in {"built-in", "auto"} and not office_pdf and not pdf_visual and not legacy_modern:
        raise ValueError("An external rendering engine is unavailable for this conversion.")
    if fmt not in {"txt", "csv"} and selected_engine == "built-in":
        unsupported = advanced_content(source)
        if unsupported:
            raise ValueError("This file contains advanced Office content that native conversion would lose: "
                             + ", ".join(unsupported[:3]) + ". Use its original editor for an exact export.")
    if office_pdf and selected_engine != "built-in":
        export_pdf(source, target, selected_engine)
        return
    if legacy_modern:
        convert_legacy(source, target, engine)
        return
    if pdf_visual:
        _pdf_to_docx_visual(source, target)
        return
    group = kind(source)
    if group == "Image":
        _save_image(source, target, fmt)
    elif group == "PDF" and fmt == "txt":
        target.write_text(_pdf_text(source), encoding="utf-8")
    elif group == "PDF" and fmt == "docx":
        document = Document()
        pages = PdfReader(str(source)).pages
        for index, page in enumerate(pages):
            if index:
                document.add_page_break()
            for line in (page.extract_text() or "").splitlines():
                document.add_paragraph(line)
        document.save(target)
    elif group == "PDF":
        _pdf_to_image(source, target, fmt)
    elif group in {"Video", "Audio"}:
        command = [ffmpeg_binary(), "-hide_banner", "-loglevel", "error", "-nostdin", "-i", str(source), "-y"]
        if fmt in {"mp4", "mov"}:
            command += ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac"]
        elif fmt in {"webm"}:
            command += ["-c:v", "libvpx-vp9", "-c:a", "libopus"]
        elif fmt == "gif":
            command += ["-vf", "fps=12,scale=720:-1:flags=lanczos"]
        command.append(str(target))
        _run(command)
    elif source.suffix.lower() == ".csv" and fmt == "xlsx":
        _csv_to_xlsx(source, target)
    elif source.suffix.lower() == ".xlsx" and fmt == "csv":
        _xlsx_to_csv(source, target)
    elif source.suffix.lower() == ".docx" and fmt == "txt":
        document = Document(source)
        target.write_text("\n".join(p.text for p in document.paragraphs), encoding="utf-8")
    elif source.suffix.lower() == ".csv" and fmt == "txt":
        shutil.copyfile(source, target)
    elif source.suffix.lower() == ".txt" and fmt == "pdf":
        _text_to_pdf(source, target)
    elif source.suffix.lower() == ".txt" and fmt == "docx":
        _text_to_docx(source, target)
    else:
        convert_office(source, target)


def compress(paths: list[Path], target: Path, progress=None) -> None:
    """Create a lossless ZIP archive, refusing collisions and input replacement."""
    sources = [Path(path).resolve(strict=True) for path in paths]
    target = Path(target).resolve()
    if not sources or any(not path.is_file() for path in sources):
        raise ValueError("Add one or more files before compressing.")
    if len({path.name.casefold() for path in sources}) != len(sources):
        raise ValueError("Files with the same name cannot share one ZIP archive.")
    if target.exists() or target in sources:
        raise FileExistsError(f"Output already exists: {target.name}")
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        with ZipFile(target, "x", compression=ZIP_DEFLATED, compresslevel=6) as archive:
            for index, path in enumerate(sources):
                archive.write(path, arcname=path.name)
                if progress:
                    progress(index, path.name)
    except Exception:
        target.unlink(missing_ok=True)
        raise


def reduced_extension(source: Path) -> str | None:
    group = kind(source)
    if group == "Image":
        return "webp"
    if group == "PDF":
        return "pdf"
    if group == "Video":
        return "mp4"
    if group == "Audio":
        return "mp3"
    return None


def _media_duration(source: Path) -> float:
    result = subprocess.run([ffmpeg_binary(), "-hide_banner", "-i", str(source)], capture_output=True, text=True, timeout=30, shell=False)
    match = re.search(r"Duration:\s*(\d+):(\d+):(\d+(?:\.\d+)?)", result.stderr)
    if not match:
        raise ValueError("Could not read media duration.")
    hours, minutes, seconds = match.groups()
    duration = int(hours) * 3600 + int(minutes) * 60 + float(seconds)
    if duration <= 0:
        raise ValueError("Media duration must be greater than zero.")
    return duration


def trim_media(source: Path, target: Path, start: float, end: float, precise: bool = True) -> None:
    source, target = Path(source).resolve(strict=True), Path(target).resolve()
    if kind(source) not in {"Video", "Audio"} or target.suffix.lower() != source.suffix.lower():
        raise ValueError("Choose audio or video and keep its file type.")
    if not all(math.isfinite(value) for value in (start, end)) or start < 0 or end <= start:
        raise ValueError("End time must be after start time.")
    if end > _media_duration(source) + 0.1:
        raise ValueError("End time is past the end of the file.")
    if target.exists() or target == source:
        raise FileExistsError(f"Output already exists: {target.name}")
    command = [ffmpeg_binary(), "-hide_banner", "-loglevel", "error", "-nostdin", "-ss", str(start), "-i", str(source), "-t", str(end - start)]
    if not precise:
        command += ["-c", "copy"]
    with tempfile.NamedTemporaryFile(prefix=".trim-", suffix=target.suffix, dir=target.parent, delete=False) as temporary:
        pending = Path(temporary.name)
    pending.unlink()
    try:
        _run(command + ["-n", str(pending)])
        os.link(pending, target)
    finally:
        pending.unlink(missing_ok=True)


def zip_entries(source: Path) -> list[tuple[str, int, int]]:
    """Read the archive listing and reject entries unsafe to extract."""
    if source.suffix.lower() != ".zip":
        raise ValueError("Choose a ZIP archive.")
    entries = []
    names = set()
    total = 0
    with ZipFile(source) as archive:
        for info in archive.infolist():
            name = info.filename.replace("\\", "/")
            parts = name.rstrip("/").split("/")
            mode = info.external_attr >> 16
            if (not name or name.startswith("/") or any(part in {"", ".", ".."} or ":" in part
                    or part.endswith((" ", ".")) or part.split(".")[0].upper() in
                    {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}
                    for part in parts) or stat.S_ISLNK(mode) or info.flag_bits & 1):
                raise ValueError(f"Unsafe ZIP entry: {info.filename}")
            key = name.rstrip("/").casefold()
            if key in names:
                raise ValueError(f"Duplicate ZIP entry: {info.filename}")
            names.add(key)
            total += info.file_size
            # ponytail: fixed archive limits; raise them if large trusted archives become a user need.
            if len(names) > 10000 or total > 2 * 1024**3 or info.file_size > 1000 * max(info.compress_size, 1):
                raise ValueError("ZIP archive exceeds safe extraction limits.")
            entries.append((name, info.file_size, info.compress_size))
    return entries


def extract_zip(source: Path, folder: Path) -> None:
    source, folder = Path(source).resolve(strict=True), Path(folder).resolve()
    zip_entries(source)
    folder.mkdir(parents=True, exist_ok=False)
    total = 0
    try:
        with ZipFile(source) as archive:
            for info in archive.infolist():
                name = info.filename.replace("\\", "/")
                output = folder.joinpath(*name.rstrip("/").split("/"))
                if not output.resolve().is_relative_to(folder):
                    raise ValueError(f"Unsafe ZIP entry: {name}")
                if info.is_dir():
                    output.mkdir(parents=True, exist_ok=True)
                    continue
                output.parent.mkdir(parents=True, exist_ok=True)
                with archive.open(info) as stream, output.open("xb") as target:
                    while block := stream.read(1024 * 1024):
                        total += len(block)
                        if total > 2 * 1024**3:
                            raise ValueError("ZIP archive exceeds safe extraction limits.")
                        target.write(block)
    except Exception:
        shutil.rmtree(folder)
        raise


OCR_LANGUAGES = {
    "multilingual": None,  # PP-OCRv6: Latin, Chinese, and Japanese scripts.
    "cyrillic": "cyrillic_PP-OCRv5_rec_mobile.onnx",
    "hindi": "devanagari_PP-OCRv5_rec_mobile.onnx",
    "korean": "korean_PP-OCRv5_rec_mobile.onnx",
    "arabic": "arabic_PP-OCRv5_rec_mobile.onnx",
}


def ocr_to_text(source: Path, target: Path, language: str = "multilingual") -> None:
    source, target = Path(source).resolve(strict=True), Path(target).resolve()
    if kind(source) not in {"Image", "PDF"}:
        raise ValueError("OCR supports images and PDFs.")
    if target.exists() or target == source:
        raise FileExistsError(f"Output already exists: {target.name}")
    if language not in OCR_LANGUAGES:
        raise ValueError(f"Unsupported OCR language: {language}")
    model_dir = Path(rapidocr.__file__).resolve().parent / "models"
    models = {"Det.model_path": "PP-OCRv6_det_small.onnx",
              "Cls.model_path": "ch_ppocr_mobile_v2.0_cls_mobile.onnx",
              "Rec.model_path": "PP-OCRv6_rec_small.onnx"}
    if any(not (model_dir / name).is_file() for name in models.values()):
        raise RuntimeError("Bundled OCR models are missing. Reinstall the app.")
    params = {"Rec.lang_type": "en", **{key: str(model_dir / name) for key, name in models.items()}}
    if language != "multilingual":
        recognition = Path(__file__).resolve().parent / "ocr_models" / OCR_LANGUAGES[language]
        if not recognition.is_file():
            raise RuntimeError(f"Bundled {language} OCR model is missing. Reinstall the app.")
        params.update({"Rec.lang_type": "devanagari" if language == "hindi" else language,
                       "Rec.ocr_version": OCRVersion.PPOCRV5,
                       "Rec.model_type": ModelType.MOBILE,
                       "Rec.model_path": str(recognition)})
    engine = RapidOCR(params=params)
    text = []
    if kind(source) == "PDF":
        pdf = pdfium.PdfDocument(str(source))
        try:
            for index in range(len(pdf)):
                page = pdf[index]
                try:
                    image = page.render(scale=2).to_pil()
                    lines = engine(image).txts or ()
                    text.append("\n".join(_ocr_lines(lines, language)))
                finally:
                    page.close()
        finally:
            pdf.close()
    else:
        with _image(source) as opened:
            image = ImageOps.exif_transpose(opened).convert("RGB")
            lines = engine(image).txts or ()
            text.append("\n".join(_ocr_lines(lines, language)))
    with target.open("x", encoding="utf-8") as output:
        output.write("\n\n".join(text) + "\n")


def _ocr_lines(lines: tuple[str, ...], language: str) -> tuple[str, ...]:
    if language == "arabic":
        # RapidOCR returns visual-order Arabic for display; text files need logical order.
        from bidi.algorithm import get_display
        return tuple(get_display(line) for line in lines)
    return lines


def reduce_size(source: Path, target: Path, target_mb: float) -> None:
    """Try to meet a per-file upload limit without truncating media."""
    source, target = Path(source).resolve(strict=True), Path(target).resolve()
    limit = int(target_mb * 1024 * 1024)
    if limit <= 0 or not source.is_file():
        raise ValueError("Choose a valid file and a target size above 0 MB.")
    if source.stat().st_size <= limit:
        raise ValueError("This file is already below the target size.")
    if target.exists() or source == target:
        raise FileExistsError(f"Output already exists: {target.name}")
    expected = reduced_extension(source)
    if expected is None or target.suffix.lower() != f".{expected}":
        raise ValueError("Target-size compression is supported for images, PDFs, video, and audio.")
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        if kind(source) == "Image":
            with _image(source) as opened:
                if getattr(opened, "n_frames", 1) > 1:
                    raise ValueError("Animated images need ZIP compression or a manual conversion.")
                image = ImageOps.exif_transpose(opened).convert("RGBA")
            for scale in (1, 0.85, 0.7, 0.55):
                resized = image if scale == 1 else image.resize((max(1, int(image.width * scale)), max(1, int(image.height * scale))), Image.Resampling.LANCZOS)
                for quality in (85, 65, 45, 25):
                    resized.save(target, "WEBP", quality=quality, method=6)
                    if target.stat().st_size <= limit:
                        return
        elif kind(source) == "PDF":
            for scale, quality in ((2, 80), (1.6, 65), (1.3, 50)):
                pdf = pdfium.PdfDocument(str(source))
                document = canvas.Canvas(str(target), pageCompression=1)
                try:
                    for page in pdf:
                        width, height = page.get_size()
                        image = page.render(scale=scale).to_pil().convert("RGB")
                        buffer = io.BytesIO()
                        image.save(buffer, "JPEG", quality=quality, optimize=True)
                        buffer.seek(0)
                        document.setPageSize((width, height))
                        document.drawImage(ImageReader(buffer), 0, 0, width, height)
                        document.showPage()
                    document.save()
                finally:
                    pdf.close()
                if target.stat().st_size <= limit:
                    return
        else:
            duration = _media_duration(source)
            is_video = kind(source) == "Video"
            for factor in (0.9, 0.75, 0.6):
                total_bitrate = max(24_000, int(limit * 8 / duration * factor))
                command = [ffmpeg_binary(), "-hide_banner", "-loglevel", "error", "-nostdin", "-i", str(source), "-y"]
                if is_video:
                    audio_bitrate = min(96_000, max(24_000, total_bitrate // 4))
                    command += ["-c:v", "libx264", "-b:v", str(max(20_000, total_bitrate - audio_bitrate)), "-c:a", "aac", "-b:a", str(audio_bitrate), "-pix_fmt", "yuv420p"]
                else:
                    command += ["-vn", "-c:a", "libmp3lame", "-b:a", str(min(320_000, total_bitrate))]
                _run(command + [str(target)])
                if target.stat().st_size <= limit:
                    return
        raise ValueError("Could not reach the target size while keeping the full file. Try a larger limit.")
    except Exception:
        target.unlink(missing_ok=True)
        raise
