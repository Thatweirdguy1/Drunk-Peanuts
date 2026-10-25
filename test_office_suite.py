"""Round trips for the native office workspaces using temporary files."""

import os
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication, QTableWidgetItem
from PySide6.QtCore import QCoreApplication, QEvent
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches
from zipfile import ZipFile
from office_suite import Dazzle, Organize, Pages, Sketch, Tables, quote


class OfficeSuiteTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def tearDown(self):
        # Dispose native widgets while QApplication is alive, not during Python shutdown.
        for widget in self.app.topLevelWidgets():
            widget.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)

    def test_roundtrips(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            pages = Pages()
            pages.editor.setPlainText("Hello office")
            pages._write(root / "note.docx")
            pages.editor.clear()
            pages.path = None
            self.assertEqual(__import__("docx").Document(root / "note.docx").paragraphs[0].text, "Hello office")

            tables = Tables()
            tables.grid.setItem(0, 0, QTableWidgetItem("42"))
            tables._write(root / "table.xlsx")
            tables.new()
            tables._read(root / "table.xlsx")
            self.assertEqual(tables.grid.item(0, 0).text(), "42")

            dazzle = Dazzle()
            dazzle.title.setPlainText("Title")
            dazzle.body.setPlainText("Body")
            dazzle._write(root / "slides.pptx")
            self.assertEqual(__import__("pptx").Presentation(root / "slides.pptx").slides[0].shapes.title.text, "Title")

            sketch = Sketch()
            sketch._write(root / "image.png")
            self.assertTrue((root / "image.png").stat().st_size > 0)

            organize = Organize()
            organize.path = root / "data.sqlite"
            with closing(sqlite3.connect(organize.path)) as db:
                with db:
                    db.execute("CREATE TABLE _freeconvert_meta (key TEXT PRIMARY KEY, value TEXT)")
                    db.execute("INSERT INTO _freeconvert_meta VALUES ('owner', 'Organize')")
                    db.execute("CREATE TABLE notes (title TEXT)")
            organize.refresh()
            organize.tables.setCurrentRow(0)
            organize.grid.setRowCount(1)
            organize.grid.setItem(0, 0, QTableWidgetItem("Local"))
            organize.save()
            with closing(sqlite3.connect(organize.path)) as db:
                self.assertEqual(db.execute("SELECT title FROM notes").fetchone()[0], "Local")
            external = root / "external.sqlite"
            with closing(sqlite3.connect(external)) as db:
                with db:
                    db.execute("CREATE TABLE original (value TEXT)")
            organize.path = external
            self.assertFalse(organize._editable())
            self.assertEqual(quote('a"b'), '"a""b"')

    def test_tables_keep_other_worksheets(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source.xlsx"
            book = Workbook()
            book.active["A1"] = "first"
            book.create_sheet("Other")["B2"] = "keep me"
            book.save(source)
            editor = Tables()
            editor._read(source)
            editor.grid.setItem(0, 0, QTableWidgetItem("changed"))
            target = root / "changed.xlsx"
            editor._write(target)
            saved = load_workbook(target)
            self.assertEqual(saved.active["A1"].value, "changed")
            self.assertEqual(saved["Other"]["B2"].value, "keep me")
            saved.close()
            self.assertEqual(load_workbook(source).active["A1"].value, "first")

    def test_formatted_runs_and_unrelated_excel_parts_survive(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source.docx"
            document = __import__("docx").Document()
            paragraph = document.add_paragraph()
            paragraph.add_run("Bold ").bold = True
            paragraph.add_run("plain")
            document.save(source)
            pages = Pages()
            pages.imported_docx = source
            pages.original_paragraphs = ["Bold plain"]
            pages.editor.setPlainText("Bold changed")
            pages._write(root / "edited.docx")
            saved = __import__("docx").Document(root / "edited.docx")
            self.assertEqual(saved.paragraphs[0].text, "Bold changed")
            self.assertTrue(saved.paragraphs[0].runs[0].bold)
            self.assertFalse(bool(saved.paragraphs[0].runs[1].bold))

            source = root / "source.xlsx"
            book = Workbook()
            book.active["A1"] = "before"
            book.create_sheet("Other")["B2"] = "untouched"
            book.save(source)
            tables = Tables()
            tables._read(source)
            tables.grid.setItem(0, 0, QTableWidgetItem("after"))
            target = root / "edited.xlsx"
            tables._write(target)
            with ZipFile(source) as old, ZipFile(target) as new:
                self.assertEqual(set(old.namelist()), set(new.namelist()))
                changed = [name for name in old.namelist() if old.read(name) != new.read(name)]
                self.assertEqual(changed, ["xl/worksheets/sheet1.xml"])
            tables.grid.setItem(0, 0, QTableWidgetItem("again"))
            tables._write(target)
            saved = load_workbook(target)
            self.assertEqual(saved.active["A1"].value, "again")
            saved.close()

    def test_dazzle_keeps_nontext_shapes_on_text_edit(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = "Before"
            shape = slide.shapes.add_shape(1, Inches(1), Inches(2), Inches(2), Inches(1))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(255, 0, 0)
            deck.save(source)
            editor = Dazzle()
            editor.imported_pptx = source
            editor.original_slides = [("Before", "")]
            editor.slides = [("Before", "")]
            editor.title.setPlainText("After")
            output = root / "edited.pptx"
            editor._write(output)
            saved = Presentation(output)
            self.assertEqual(saved.slides[0].shapes.title.text, "After")
            self.assertEqual(len(saved.slides[0].shapes), len(slide.shapes))

    def test_dazzle_keeps_inline_text_formatting(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            paragraph = slide.shapes.title.text_frame.paragraphs[0]
            paragraph.add_run().text = "Bold "
            paragraph.runs[0].font.bold = True
            paragraph.add_run().text = "plain"
            deck.save(source)
            editor = Dazzle()
            editor.imported_pptx = source
            editor.original_slides = [("Bold plain", "")]
            editor.slides = [("Bold changed", "")]
            editor.title.setPlainText("Bold changed")
            editor._write(root / "edited.pptx")
            saved = Presentation(root / "edited.pptx")
            runs = saved.slides[0].shapes.title.text_frame.paragraphs[0].runs
            self.assertEqual("".join(run.text for run in runs), "Bold changed")
            self.assertTrue(runs[0].font.bold)

    def test_unchanged_office_imports_copy_exact_bytes(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            document = __import__("docx").Document()
            document.add_paragraph("Original")
            docx = root / "source.docx"
            document.save(docx)
            pages = Pages()
            pages.imported_docx = docx
            pages.original_paragraphs = ["Original"]
            pages.editor.setPlainText("Original")
            pages_copy = root / "copy.docx"
            pages._write(pages_copy)
            self.assertEqual(docx.read_bytes(), pages_copy.read_bytes())

            book = Workbook()
            book.active["A1"] = "Original"
            book.create_sheet("Extra")["A1"] = "Keep"
            xlsx = root / "source.xlsx"
            book.save(xlsx)
            tables = Tables()
            tables._read(xlsx)
            tables_copy = root / "copy.xlsx"
            tables._write(tables_copy)
            self.assertEqual(xlsx.read_bytes(), tables_copy.read_bytes())

            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = "Original"
            pptx = root / "source.pptx"
            deck.save(pptx)
            dazzle = Dazzle()
            dazzle.imported_pptx = pptx
            dazzle.original_slides = [("Original", "")]
            dazzle.slides = [("Original", "")]
            dazzle.title.setPlainText("Original")
            dazzle_copy = root / "copy.pptx"
            dazzle._write(dazzle_copy)
            self.assertEqual(pptx.read_bytes(), dazzle_copy.read_bytes())


if __name__ == "__main__":
    unittest.main()
