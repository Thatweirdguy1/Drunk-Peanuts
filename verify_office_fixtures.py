"""Check text edits on copies of the user's Office fixtures."""

from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from office_surgery import edit_xlsx_cells, replace_runs, replace_text_frame


def check_parts(source, target, allowed):
    with ZipFile(source) as old, ZipFile(target) as new:
        if set(old.namelist()) != set(new.namelist()):
            raise AssertionError(f"Package parts changed: {source.name}")
        unexpected = [name for name in old.namelist() if name not in allowed and old.read(name) != new.read(name)]
        if unexpected:
            raise AssertionError(f"Unexpected parts changed in {source.name}: {unexpected}")


def main():
    fixtures = Path("test docs")
    with TemporaryDirectory() as folder:
        for source in sorted(fixtures.glob("*.docx")):
            document = Document(source)
            paragraph_index = next((i for i, p in enumerate(document.paragraphs) if p.text and p.runs), None)
            if paragraph_index is None:
                raise AssertionError(f"No editable paragraph: {source.name}")
            paragraph = document.paragraphs[paragraph_index]
            replace_runs(paragraph, paragraph.text + " verified")
            target = Path(folder) / source.name
            document.save(target)
            check_parts(source, target, {"word/document.xml"})
            assert "verified" in Document(target).paragraphs[paragraph_index].text
            print("PASS", source.name)
        for source in sorted(fixtures.glob("*.xlsx")):
            book = load_workbook(source, read_only=True, data_only=False)
            first = book.active.cell(1, 1).value
            book.close()
            before = [["" if first is None else str(first)]]
            after = [[before[0][0] + " verified"]]
            target = Path(folder) / source.name
            edit_xlsx_cells(source, target, before, after)
            with ZipFile(source) as archive:
                sheet = __import__("office_surgery")._active_sheet(archive)
            check_parts(source, target, {sheet})
            book = load_workbook(target, read_only=True)
            assert str(book.active.cell(1, 1).value).endswith(" verified")
            book.close()
            print("PASS", source.name)
        for source in sorted(fixtures.glob("*.pptx")):
            deck = Presentation(source)
            slide = next((slide for slide in deck.slides if any(s.has_text_frame and s.text.strip() for s in slide.shapes)), None)
            if slide is None:
                raise AssertionError(f"No editable slide: {source.name}")
            shape = next(s for s in slide.shapes if s.has_text_frame and s.text.strip())
            replace_text_frame(shape, shape.text + " verified")
            target = Path(folder) / source.name
            deck.save(target)
            slide_index = list(deck.slides).index(slide) + 1
            check_parts(source, target, {f"ppt/slides/slide{slide_index}.xml"})
            assert any(s.text.endswith(" verified") for s in Presentation(target).slides[slide_index - 1].shapes if s.has_text_frame)
            print("PASS", source.name)


if __name__ == "__main__":
    main()
