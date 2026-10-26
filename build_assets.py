"""Create original Drunk Peanuts branding and editable local Office starters."""
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
import sqlite3
from html import escape
from contextlib import closing
from office_templates import style_slide


def build_branding(folder=Path('assets')):
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    logo=Image.new('RGB',(1600,320),'black');draw=ImageDraw.Draw(logo)
    font=ImageFont.truetype('C:/Windows/Fonts/BRUSHSCI.TTF',164)
    draw.text((800,150),'Drunk Peanuts.',font=font,fill='#ffdd22',anchor='mm',stroke_width=2)
    draw.rounded_rectangle((400,265,1200,278),radius=6,fill='#a964eb');logo.save(folder/'logo.png')
    tile=Image.new('RGB',(256,256),'black');draw=ImageDraw.Draw(tile);font=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',125)
    draw.rounded_rectangle((10,10,246,246),radius=35,outline='#a964eb',width=12)
    draw.text((73,130),'D',font=font,fill='#a964eb',anchor='mm');draw.text((172,130),'P',font=font,fill='#ffdd22',anchor='mm')
    tile.save(folder/'app.icns')
    tile.save(folder/'app.ico',sizes=[(16,16),(24,24),(32,32),(48,48),(64,64),(128,128),(256,256)])


def build(folder=Path('assets')):
    folder=Path(folder);build_branding(folder)
    templates = folder/'templates'
    templates.mkdir(exist_ok=True)
    for name, lines in {
        'Report': [('Title','Report title'), ('Subtitle','Author · Date'), ('Heading 1','Summary'), ('Normal','Write a brief overview.'), ('Heading 1','Findings'), ('Normal','Add your findings here.'), ('Heading 1','Next steps'), ('Normal','List actions and owners.')],
        'Letter': [('Title','Your name'), ('Normal','Address · Email · Phone'), ('Normal','Date'), ('Normal','Recipient name\nRecipient address'), ('Normal','Dear recipient,'), ('Normal','Write your letter here.'), ('Normal','Sincerely,\nYour name')],
        'Meeting notes': [('Title','Meeting notes'), ('Subtitle','Date · Attendees'), ('Heading 1','Agenda'), ('Normal','1. First topic\n2. Second topic'), ('Heading 1','Decisions'), ('Normal','Record decisions here.'), ('Heading 1','Action items'), ('Normal','Task · Owner · Due date')],
    }.items():
        doc = Document()
        for style, text in lines:
            doc.add_paragraph(text, style)
        doc.save(templates/f'{name}.docx')
        tags = {'Title': 'h1', 'Subtitle': 'h3', 'Heading 1': 'h2', 'Normal': 'p'}
        html = ''.join(f'<{tags[style]}>{escape(text).replace(chr(10), "<br>")}</{tags[style]}>' for style, text in lines)
        (templates/f'{name}.html').write_text('<html><head><style>body {font-family: Calibri; font-size:12pt;} h1 {font-size:28pt;} h2 {font-size:17pt;} h3 {font-size:13pt; color:#64748b;} p {margin-top:8pt; margin-bottom:8pt;}</style></head><body>'+html+'</body></html>',encoding='utf-8')
    for name, rows in {
        'Budget': [['Category','Planned','Actual','Difference'],['Housing',0,0,'=B2-C2'],['Food',0,0,'=B3-C3'],['Transport',0,0,'=B4-C4'],['Total','=SUM(B2:B4)','=SUM(C2:C4)','=B5-C5']],
        'Task tracker': [['Task','Owner','Due date','Status'],['First task','Your name','YYYY-MM-DD','To do']],
    }.items():
        book=Workbook(); sheet=book.active; sheet.title=name
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        for col in ('A','B','C','D'):
            sheet.column_dimensions[col].width=24
        book.save(templates/f'{name}.xlsx')
    decks = {
        'Presentation': [('Presentation title','Subtitle · Presenter'),('Agenda','Context\nKey points\nNext steps'),('Key points','Add your main message here.'),('Next steps','Action · Owner · Date')],
        'Project update': [('Project update','Project · Date'),('Progress','Completed\nIn progress'),('Risks','Risk · Impact · Mitigation'),('Next milestone','Deliverable · Owner · Date')],
        'Pitch deck': [('Your big idea','Company · Presenter · Date'),('The problem','Who experiences it?\nWhy does it matter?'),('Our solution','What we offer\nHow it solves the problem'),('Audience and opportunity','Target audience\nEvidence and market context'),('Plan and milestones','Milestone · Owner · Date\nResources required'),('The ask','Support needed\nNext conversation')],
        'Research presentation': [('Research title','Author · Institution · Date'),('Research question','Question\nWhy it matters'),('Method','Participants or data\nProcedure\nAnalysis'),('Results','Key finding\nSupporting evidence'),('Discussion','Interpretation\nLimitations'),('Conclusion','Takeaway\nFuture work\nReferences')],
        'Lesson plan': [('Lesson title','Subject · Level · Teacher'),('Learning objectives','By the end, learners can…\nPrior knowledge'),('Explore the topic','Key concept\nWorked example'),('Practice','Activity instructions\nTime and materials'),('Review and reflect','Check understanding\nKey takeaway\nNext lesson')],
        'Portfolio': [('Your name','Role · Speciality · Contact'),('About me','Background\nWhat I bring'),('Selected work','Project title\nRole and contribution\nOutcome'),('Case study','Challenge\nApproach\nResult'),('Let’s connect','Skills\nAvailability\nContact details')],
    }
    for name, slides in decks.items():
        deck=Presentation()
        deck.slide_width=Inches(13.333333)
        deck.slide_height=Inches(7.5)
        for title, body in slides:
            slide=deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text=title; slide.placeholders[1].text=body
            style_slide(slide)
        deck.save(templates/f'{name}.pptx')
    (templates/'Equations.formulaic').write_text('x = {-b + sqrt {b^2 - 4 a c}} over {2 a}', encoding='utf-8')
    Image.new('RGB',(1600,900),'white').save(templates/'Blank canvas.png')
    (templates/'Screenplay.fountain').write_text('Title: Untitled screenplay\nAuthor: Your name\n\nINT. LOCATION - DAY\n\nDescribe the scene.\n\nCHARACTER\nDialogue goes here.\n',encoding='utf-8')
    database=templates/'Contacts.sqlite'
    if not database.exists():
        with sqlite3.connect(database) as db:
            db.execute('CREATE TABLE Contacts (Name TEXT, Email TEXT, Phone TEXT, Notes TEXT)')
            db.execute('INSERT INTO Contacts VALUES (?, ?, ?, ?)',('Example contact','name@example.com','','Replace this sample'))
    with closing(sqlite3.connect(database)) as db:
        with db:
            db.execute('CREATE TABLE IF NOT EXISTS _freeconvert_meta (key TEXT PRIMARY KEY, value TEXT)')
            db.execute("INSERT OR REPLACE INTO _freeconvert_meta VALUES ('owner', 'Organize')")

if __name__ == '__main__':
    build()
