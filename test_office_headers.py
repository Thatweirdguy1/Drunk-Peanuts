import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import io,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from PySide6.QtCore import QCoreApplication,QEvent
from PySide6.QtWidgets import QApplication
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader
from PIL import Image
from office_suite import Pages
from office_headers import set_headers,validate


class HeaderFooterTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def tearDown(self):
        for widget in self.app.topLevelWidgets():widget.deleteLater()
        QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete)
    def settings(self,**overrides):return dict(dict(header='Drunk Peanuts report',footer='Local and private',numbers='Footer',align='Right',start=4,skip_first=True),**overrides)

    def test_native_fields_first_page_and_existing_body_preserved(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);source=root/'source.docx';target=root/'result.docx';doc=Document();paragraph=doc.add_paragraph('Keep this original field and image ')
            run=paragraph.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');code=OxmlElement('w:instrText');code.text='DATE';end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');run._r.extend([begin,code,end]);image=io.BytesIO();Image.new('RGB',(10,10),'purple').save(image,format='PNG');image.seek(0);doc.add_picture(image);doc.save(source)
            original=source.read_bytes();body=[p.xml for p in doc.element.body if p.tag!=qn('w:sectPr')];pages=Pages();pages.open(source);set_headers(pages,self.settings());pages._write(target)
            saved=Document(target);self.assertEqual(body,[p.xml for p in saved.element.body if p.tag!=qn('w:sectPr')]);self.assertEqual(source.read_bytes(),original)
            section=saved.sections[0];self.assertTrue(section.different_first_page_header_footer);self.assertEqual(section._sectPr.find(qn('w:pgNumType')).get(qn('w:start')),'4');self.assertIn('Drunk Peanuts report',section.header.paragraphs[0].text)
            fields=section.footer._element.xpath('.//w:instrText');self.assertEqual([n.text.strip() for n in fields],['PAGE','NUMPAGES']);self.assertEqual(section.first_page_footer.paragraphs[0].text,'')
            pages.new();self.assertIsNone(pages.header_footer)

    def test_pdf_page_numbers_correct_skip_first_and_plain_export_guard(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);pages=Pages();pages.editor.setHtml('<p>First page</p><p style="page-break-before:always">Second page</p><p style="page-break-before:always">Third page</p>');set_headers(pages,self.settings(start=1));target=root/'numbered.pdf'
            with patch('office_suite.pick_save',return_value=target):pages.export_pdf()
            reader=PdfReader(target);self.assertEqual(len(reader.pages),3);self.assertNotIn('Local and private',reader.pages[0].extract_text());self.assertIn('Page 2 of 3',reader.pages[1].extract_text());self.assertIn('Page 3 of 3',reader.pages[2].extract_text());self.assertIn('Drunk Peanuts report',reader.pages[2].extract_text())
            with self.assertRaises(ValueError):pages._write(root/'loses-headers.txt')
            with self.assertRaises(ValueError):validate(self.settings(start=0))


if __name__=='__main__':unittest.main()
