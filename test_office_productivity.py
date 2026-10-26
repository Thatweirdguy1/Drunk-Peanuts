import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import tempfile,unittest
from pathlib import Path
from zipfile import ZipFile
from PySide6.QtCore import Qt,QCoreApplication,QEvent
from PySide6.QtWidgets import QApplication
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from office_suite import Pages
from office_productivity import set_properties,statistics,update_outline
from office_ribbon import install_ribbon


class ProductivityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def tearDown(self):
        for widget in self.app.topLevelWidgets():widget.deleteLater()
        QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete)

    def test_properties_preserve_native_body_fields_and_other_properties(self):
        with tempfile.TemporaryDirectory() as folder:
            source=Path(folder)/'source.docx';target=Path(folder)/'edited.docx'
            doc=Document();doc.add_heading('Introduction',1);run=doc.add_paragraph('Original body ').add_run()
            field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'DATE');run._r.append(field)
            doc.core_properties.author='Original author';doc.core_properties.language='en-IN';doc.save(source)
            body=Document(source).element.body.xml;pages=Pages();pages.open(source)
            self.assertEqual(pages.document_properties['author'],'Original author')
            set_properties(pages,{'title':'Quarterly report','keywords':'local, offline'})
            self.assertTrue(pages.editor.document().isModified())
            with self.assertRaises(ValueError):pages._write(Path(folder)/'wrong.txt')
            pages._write(target);result=Document(target)
            self.assertEqual(result.element.body.xml,body)
            self.assertEqual(result.core_properties.title,'Quarterly report')
            self.assertEqual(result.core_properties.author,'Original author')
            self.assertEqual(result.core_properties.language,'en-IN')
            self.assertEqual(Document(source).core_properties.title,'')
            self.assertFalse(pages.pending_properties);self.assertFalse(pages.editor.document().isModified())
            pages.open(target);self.assertEqual(pages.document_properties['title'],'Quarterly report')
            with self.assertRaises(ValueError):set_properties(pages,{'author':'a'*256})
            pages.new();self.assertEqual(pages.document_properties['author'],'')

    def test_outline_tracks_headings_filters_and_jumps_without_editing(self):
        pages=Pages();pages.editor.setHtml('<h1>Overview</h1><p>Body text</p><h2>Details</h2><p>More</p>')
        pages.editor.document().setModified(False);install_ribbon(pages);update_outline(pages)
        self.assertEqual(pages.outline_list.count(),2)
        item=pages.outline_list.item(1);pages.outline_list.itemActivated.emit(item)
        self.assertEqual(pages.editor.textCursor().block().text(),'Details')
        self.assertFalse(pages.editor.document().isModified())
        pages.outline_search.setText('OVER');self.assertEqual(pages.outline_list.count(),1)
        pages.outline_search.clear();pages.editor.setHtml('<h1>New section</h1>');update_outline(pages)
        self.assertEqual(pages.outline_list.item(0).text(),'New section')
        self.assertIn('Document properties',pages.ribbon_buttons)
        self.assertIn('Document statistics',pages.ribbon_buttons)
        self.assertEqual(statistics("Hello, world!\nIt's local.\n\ufffc")['words'],4)
        self.assertEqual(statistics('One two\n\nThree')['nonempty_paragraphs'],2)


if __name__=='__main__':unittest.main()
