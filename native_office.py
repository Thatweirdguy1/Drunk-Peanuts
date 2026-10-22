"""Local Office format conversions without a separate office process."""

from __future__ import annotations

import csv
from html import escape
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from odf import draw, style, table, text
from odf.opendocument import (OpenDocumentPresentation, OpenDocumentSpreadsheet,
                              OpenDocumentText, load as load_odf)
from odf.teletype import extractText
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from striprtf.striprtf import rtf_to_text


DOCUMENTS = {"txt", "docx", "odt", "rtf"}
SPREADSHEETS = {"csv", "xlsx", "ods"}
PRESENTATIONS = {"pptx", "odp"}


def targets(extension: str) -> set[str]:
    if extension in DOCUMENTS:
        return (DOCUMENTS | {"pdf"}) - {extension}
    if extension in SPREADSHEETS:
        return (SPREADSHEETS | {"pdf"}) - {extension}
    if extension in PRESENTATIONS:
        return (PRESENTATIONS | {"pdf"}) - {extension}
    return set()


def _docx_blocks(source: Path) -> list[dict]:
    document = Document(source)
    blocks = []
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            paragraph = DocxParagraph(element, document)
            name = paragraph.style.name if paragraph.style else ""
            blocks.append({"kind": "p", "heading": name.startswith("Heading"),
                           "runs": [(run.text, bool(run.bold), bool(run.italic), bool(run.underline)) for run in paragraph.runs]})
        elif element.tag == qn("w:tbl"):
            grid = DocxTable(element, document)
            blocks.append({"kind": "table", "rows": [[cell.text for cell in row.cells] for row in grid.rows]})
    return blocks


def _odt_blocks(source: Path) -> list[dict]:
    document = load_odf(str(source))
    blocks = []
    for element in document.text.childNodes:
        if element.qname == text.P().qname or element.qname == text.H(outlinelevel=1).qname:
            blocks.append({"kind": "p", "heading": element.qname == text.H(outlinelevel=1).qname,
                           "runs": [(extractText(element), False, False, False)]})
        elif element.qname == table.Table(name="x").qname:
            rows = []
            for row in element.getElementsByType(table.TableRow):
                values = [extractText(cell) for cell in row.getElementsByType(table.TableCell)]
                if values:
                    rows.append(values)
            blocks.append({"kind": "table", "rows": rows})
    return blocks


def read_document(source: Path) -> list[dict]:
    ext = source.suffix.lower().lstrip(".")
    if ext == "docx":
        return _docx_blocks(source)
    if ext == "odt":
        return _odt_blocks(source)
    if ext == "rtf":
        content = rtf_to_text(source.read_text(encoding="utf-8", errors="replace"))
    elif ext == "txt":
        content = source.read_text(encoding="utf-8-sig")
    else:
        raise ValueError("Unsupported document format")
    return [{"kind": "p", "heading": False, "runs": [(line, False, False, False)]}
            for line in content.splitlines()]


def _plain(blocks: list[dict]) -> str:
    lines = []
    for block in blocks:
        if block["kind"] == "p":
            lines.append("".join(run[0] for run in block["runs"]))
        else:
            lines.extend("\t".join(row) for row in block["rows"])
    return "\n".join(lines)


def _write_docx(blocks: list[dict], target: Path) -> None:
    document = Document()
    for block in blocks:
        if block["kind"] == "table":
            rows = block["rows"]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            grid = document.add_table(rows=len(rows), cols=width)
            for i, row in enumerate(rows):
                for j, value in enumerate(row):
                    grid.cell(i, j).text = value
        else:
            paragraph = document.add_paragraph(style="Heading 1" if block["heading"] else None)
            for value, bold, italic, underline in block["runs"]:
                run = paragraph.add_run(value)
                run.bold, run.italic, run.underline = bold, italic, underline
    document.save(target)


def _write_odt(blocks: list[dict], target: Path) -> None:
    document = OpenDocumentText()
    for block in blocks:
        if block["kind"] == "table":
            grid = table.Table(name=f"Table{len(document.text.getElementsByType(table.Table)) + 1}")
            for values in block["rows"]:
                row = table.TableRow()
                for value in values:
                    cell = table.TableCell(valuetype="string")
                    cell.addElement(text.P(text=value))
                    row.addElement(cell)
                grid.addElement(row)
            document.text.addElement(grid)
        else:
            content = "".join(run[0] for run in block["runs"])
            paragraph = text.H(outlinelevel=1, text=content) if block["heading"] else text.P(text=content)
            document.text.addElement(paragraph)
    document.save(str(target))


def _write_rtf(blocks: list[dict], target: Path) -> None:
    def encode(value):
        output = []
        for character in value:
            if character in "{}\\":
                output.append("\\" + character)
            elif ord(character) < 128:
                output.append(character)
            else:
                codepoint = ord(character)
                if codepoint > 65535:
                    codepoint -= 65536
                    units = [0xD800 + (codepoint >> 10), 0xDC00 + (codepoint & 1023)]
                else:
                    units = [codepoint]
                output.extend(f"\\u{unit if unit < 32768 else unit - 65536}?" for unit in units)
        return "".join(output)
    lines = ["{\\rtf1\\ansi\\deff0"]
    for block in blocks:
        if block["kind"] == "table":
            lines.extend(encode("\t".join(row)) + r"\par" for row in block["rows"])
        else:
            line = ""
            for value, bold, italic, underline in block["runs"]:
                line += (r"\b " if bold else "") + (r"\i " if italic else "") + (r"\ul " if underline else "")
                line += encode(value)
                line += (r"\ul0 " if underline else "") + (r"\i0 " if italic else "") + (r"\b0 " if bold else "")
            lines.append(line + r"\par")
    target.write_text("\n".join(lines) + "}", encoding="ascii")


def _write_document_pdf(blocks: list[dict], target: Path) -> None:
    normal = ParagraphStyle("Body", fontName="Helvetica", fontSize=11, leading=15, spaceAfter=7, alignment=TA_LEFT)
    heading = ParagraphStyle("Heading", parent=normal, fontName="Helvetica-Bold", fontSize=17, leading=21, spaceBefore=12)
    story = []
    for block in blocks:
        if block["kind"] == "table":
            if block["rows"]:
                rows = [[Paragraph(escape(str(value)), normal) for value in row] for row in block["rows"]]
                width = max(len(row) for row in rows)
                rows = [row + [""] * (width - len(row)) for row in rows]
                grid = Table(rows, repeatRows=1, hAlign="LEFT")
                grid.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b8c8cc")),
                                          ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                          ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef5f4"))]))
                story.extend([grid, Spacer(1, 9)])
        else:
            markup = "".join(("<b>" if bold else "") + ("<i>" if italic else "") + ("<u>" if underline else "") + escape(value)
                             + ("</u>" if underline else "") + ("</i>" if italic else "") + ("</b>" if bold else "")
                             for value, bold, italic, underline in block["runs"])
            story.append(Paragraph(markup or "&nbsp;", heading if block["heading"] else normal))
    SimpleDocTemplate(str(target), pagesize=A4, leftMargin=48, rightMargin=48, topMargin=48, bottomMargin=48).build(story)


def convert_document(source: Path, target: Path) -> None:
    blocks = read_document(source)
    fmt = target.suffix.lower().lstrip(".")
    if fmt == "txt":
        target.write_text(_plain(blocks), encoding="utf-8")
    elif fmt == "docx":
        _write_docx(blocks, target)
    elif fmt == "odt":
        _write_odt(blocks, target)
    elif fmt == "rtf":
        _write_rtf(blocks, target)
    elif fmt == "pdf":
        _write_document_pdf(blocks, target)
    else:
        raise ValueError("Unsupported document output")


def _sheet_rows(source: Path) -> list[list[str]]:
    ext = source.suffix.lower().lstrip(".")
    if ext == "csv":
        with source.open(encoding="utf-8-sig", newline="") as stream:
            return list(csv.reader(stream))
    if ext == "xlsx":
        book = load_workbook(source, read_only=True, data_only=True)
        try:
            return [["" if cell is None else str(cell) for cell in row] for row in book.active.iter_rows(values_only=True)]
        finally:
            book.close()
    document = load_odf(str(source))
    sheet = document.spreadsheet.getElementsByType(table.Table)[0]
    rows = []
    for row in sheet.getElementsByType(table.TableRow):
        values = []
        for cell in row.getElementsByType(table.TableCell):
            value = extractText(cell) or cell.getAttribute("value") or ""
            values.extend([str(value)] * min(int(cell.getAttribute("numbercolumnsrepeated") or 1), 10000))
        rows.extend([values] * min(int(row.getAttribute("numberrowsrepeated") or 1), 10000))
    return rows


def convert_spreadsheet(source: Path, target: Path) -> None:
    rows = _sheet_rows(source)
    fmt = target.suffix.lower().lstrip(".")
    if fmt == "csv":
        with target.open("w", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream).writerows(rows)
    elif fmt == "xlsx":
        book = Workbook(write_only=True)
        sheet = book.create_sheet("Sheet1")
        for row in rows:
            sheet.append(row)
        book.save(target)
    elif fmt == "ods":
        document = OpenDocumentSpreadsheet()
        sheet = table.Table(name="Sheet1")
        for values in rows:
            row = table.TableRow()
            for value in values:
                cell = table.TableCell(valuetype="string")
                cell.addElement(text.P(text=value))
                row.addElement(cell)
            sheet.addElement(row)
        document.spreadsheet.addElement(sheet)
        document.save(str(target))
    elif fmt == "pdf":
        normal = ParagraphStyle("Cell", fontName="Helvetica", fontSize=8, leading=10)
        printable = [[Paragraph(escape(value), normal) for value in row] for row in rows]
        if printable:
            width = max(len(row) for row in printable)
            printable = [row + [""] * (width - len(row)) for row in printable]
            grid = Table(printable, repeatRows=1, hAlign="LEFT")
            grid.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.grey)]))
            SimpleDocTemplate(str(target), pagesize=landscape(A4)).build([grid])
        else:
            canvas.Canvas(str(target), pagesize=landscape(A4)).save()
    else:
        raise ValueError("Unsupported spreadsheet output")


def _slides(source: Path) -> list[list[str]]:
    if source.suffix.lower() == ".pptx":
        deck = Presentation(source)
        return [[shape.text for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
                for slide in deck.slides]
    document = load_odf(str(source))
    return [[extractText(paragraph) for paragraph in page.getElementsByType(text.P)]
            for page in document.presentation.getElementsByType(draw.Page)]


def convert_presentation(source: Path, target: Path) -> None:
    slides = _slides(source)
    fmt = target.suffix.lower().lstrip(".")
    if fmt == "pptx":
        deck = Presentation()
        for lines in slides:
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            left, top, width = 914400, 914400, deck.slide_width - 1828800
            for line in lines:
                box = slide.shapes.add_textbox(left, top, width, 685800)
                box.text = line
                top += 822960
        deck.save(target)
    elif fmt == "odp":
        document = OpenDocumentPresentation()
        layout = style.PageLayout(name="PageLayout")
        layout.addElement(style.PageLayoutProperties(pagewidth="28cm", pageheight="15.75cm"))
        document.automaticstyles.addElement(layout)
        master = style.MasterPage(name="Standard", pagelayoutname=layout)
        document.masterstyles.addElement(master)
        for index, lines in enumerate(slides, 1):
            page = draw.Page(name=f"page{index}", masterpagename=master)
            for line_index, line in enumerate(lines):
                frame = draw.Frame(width="25cm", height="1.5cm", x="1.5cm", y=f"{1.5 + line_index * 1.7}cm")
                box = draw.TextBox()
                box.addElement(text.P(text=line))
                frame.addElement(box)
                page.addElement(frame)
            document.presentation.addElement(page)
        document.save(str(target))
    elif fmt == "pdf":
        pdf = canvas.Canvas(str(target), pagesize=landscape(A4))
        page_width, page_height = landscape(A4)
        for lines in slides:
            y = page_height - 60
            for line in lines:
                for wrapped in simpleSplit(line, "Helvetica", 20, page_width - 100):
                    pdf.setFont("Helvetica", 20)
                    pdf.drawString(50, y, wrapped)
                    y -= 30
            pdf.showPage()
        pdf.save()
    else:
        raise ValueError("Unsupported presentation output")


def convert_office(source: Path, target: Path) -> None:
    ext = source.suffix.lower().lstrip(".")
    if target.suffix.lower().lstrip(".") not in targets(ext):
        raise ValueError("Unsupported Office conversion")
    if ext in DOCUMENTS:
        convert_document(source, target)
    elif ext in SPREADSHEETS:
        convert_spreadsheet(source, target)
    else:
        convert_presentation(source, target)
