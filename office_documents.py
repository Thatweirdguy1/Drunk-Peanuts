"""Rich document insertion and Word serialization, using native Qt and python-docx."""
import base64
import html
import io
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from PySide6.QtCore import QUrl, Qt
from PySide6.QtGui import QColor, QTextImageFormat, QTextCursor, QTextDocument, QTextTable, QTextTableFormat, QTextFormat, QTextCharFormat
from PySide6.QtWidgets import QFileDialog, QInputDialog, QMessageBox
from office_insert import image_bytes, ask_chart, chart_png, video_poster


def insert_picture(pages,path=None,blob=None):
    if blob is None:
        if path is None:
            name,_=QFileDialog.getOpenFileName(pages,'Insert image','','Images (*.png *.jpg *.jpeg *.webp *.bmp)');path=Path(name) if name else None
        if not path:return
        blob=image_bytes(path)
    url='data:image/png;base64,'+base64.b64encode(blob).decode()
    fmt=QTextImageFormat();fmt.setName(url);fmt.setWidth(420)
    from PySide6.QtGui import QImage
    image=QImage.fromData(blob);fmt.setHeight(420*image.height()/max(1,image.width()))
    pages.editor.textCursor().insertImage(fmt);pages.rich_changed=True
    return url

def insert_chart(pages):
    result=ask_chart(pages)
    if result:insert_picture(pages,blob=chart_png(result[1],result[0]))

def insert_table(pages,rows=None,columns=None):
    if rows is None:
        rows,ok=QInputDialog.getInt(pages,'Insert table','Rows',3,1,100)
        if not ok:return
        columns,ok=QInputDialog.getInt(pages,'Insert table','Columns',3,1,30)
        if not ok:return
    if not 1<=rows<=100 or not 1<=columns<=30:raise ValueError('Table exceeds 100 rows or 30 columns.')
    fmt=QTextTableFormat();fmt.setBorder(1);fmt.setCellPadding(5);fmt.setCellSpacing(0)
    table=pages.editor.textCursor().insertTable(rows,columns,fmt);pages.rich_changed=True;return table

def insert_video(pages,path=None):
    if path is None:
        name,_=QFileDialog.getOpenFileName(pages,'Insert local video link','','Video (*.mp4 *.mov *.webm *.avi)');path=Path(name) if name else None
    if not path:return
    path=Path(path).resolve();poster=video_poster(path);insert_picture(pages,blob=poster)
    cursor=pages.editor.textCursor();cursor.insertBlock();cursor.insertHtml('<a href="'+html.escape(path.as_uri(),quote=True)+'">Play video: '+html.escape(path.name)+'</a>');cursor.insertBlock();pages.rich_changed=True

def insert_shape(pages):
    kind,ok=QInputDialog.getItem(pages,'Insert shape','Shape',['Rectangle','Ellipse','Text box'],0,False)
    if not ok:return
    text,ok=QInputDialog.getText(pages,'Shape text','Text')
    if not ok:return
    from PySide6.QtGui import QImage,QPainter,QPen,QColor,QFont
    from PySide6.QtCore import QBuffer,QIODevice,QRectF
    image=QImage(600,180,QImage.Format_ARGB32);image.fill(Qt.transparent);p=QPainter(image);p.setRenderHint(QPainter.Antialiasing);p.setPen(QPen(QColor('#222222'),2));p.setBrush(QColor('#ffdd22'))
    if kind=='Ellipse':p.drawEllipse(QRectF(2,2,596,176))
    else:p.drawRect(QRectF(2,2,596,176))
    p.setFont(QFont('Arial',22));p.drawText(QRectF(15,10,570,160),Qt.AlignCenter|Qt.TextWordWrap,text);p.end();b=QBuffer();b.open(QIODevice.WriteOnly);image.save(b,'PNG');insert_picture(pages,blob=bytes(b.data()))

def setup_pages(pages):
    pages.rich_changed=False
    pages.docx_loaded_rich=False
    menu=pages.menuBar().addMenu('Insert')
    for label,callback in [('Image',lambda:insert_picture(pages)),('Table',lambda:insert_table(pages)),('Chart / graph',lambda:insert_chart(pages)),('Shape / text box',lambda:insert_shape(pages)),('Video link with poster',lambda:insert_video(pages)),('Page break',lambda:page_break(pages))]:
        menu.addAction(label,lambda checked=False,cb=callback:pages.report(cb))
    menu.addAction('Text color',lambda:font_color(pages))
    menu.addAction('Paragraph background',lambda:page_color(pages))
    menu.addAction('Play linked video',lambda:pages.report(lambda:play_linked_video(pages)))

def font_color(pages):
    from PySide6.QtWidgets import QColorDialog
    color=QColorDialog.getColor(QColor('#111111'),pages)
    if color.isValid():pages.editor.setTextColor(color);pages.rich_changed=True

def page_color(pages):
    from PySide6.QtWidgets import QColorDialog
    color=QColorDialog.getColor(QColor('#ffffff'),pages)
    if color.isValid():
        cursor=pages.editor.textCursor();fmt=cursor.blockFormat();fmt.setBackground(color);cursor.setBlockFormat(fmt);pages.rich_changed=True

def page_break(pages):
    cursor=pages.editor.textCursor();cursor.insertBlock();fmt=cursor.blockFormat();fmt.setPageBreakPolicy(QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore);cursor.setBlockFormat(fmt);pages.rich_changed=True

def _hyperlink(paragraph,text,url):
    node=OxmlElement('w:hyperlink');node.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True));run=OxmlElement('w:r');t=OxmlElement('w:t');t.text=text;run.append(t);node.append(run);paragraph._p.append(node)

def _block(paragraph,block,document):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    fmt=block.blockFormat();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER if fmt.alignment()&Qt.AlignHCenter else WD_ALIGN_PARAGRAPH.RIGHT if fmt.alignment()&Qt.AlignRight else WD_ALIGN_PARAGRAPH.LEFT
    if fmt.headingLevel():paragraph.style=f'Heading {min(3,fmt.headingLevel())}'
    if block.textList():paragraph.style='List Number' if block.textList().format().style().value<=-4 else 'List Bullet'
    if fmt.lineHeightType()==1 and fmt.lineHeight():paragraph.paragraph_format.line_spacing=fmt.lineHeight()/100
    if fmt.pageBreakPolicy()&QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore:paragraph.paragraph_format.page_break_before=True
    if fmt.background().style()!=Qt.NoBrush:
        node=OxmlElement('w:shd');node.set(qn('w:fill'),fmt.background().color().name()[1:]);paragraph._p.get_or_add_pPr().append(node)
    iterator=block.begin()
    while not iterator.atEnd():
        fragment=iterator.fragment()
        if fragment.isValid():
            char=fragment.charFormat()
            if char.isImageFormat():
                image=char.toImageFormat();name=image.name()
                if not name.startswith('data:image/') or ';base64,' not in name:raise ValueError('An image resource is unavailable; reinsert the image.')
                blob=base64.b64decode(name.split(',',1)[1],validate=True);paragraph.add_run().add_picture(io.BytesIO(blob),width=Inches(max(1,image.width())/96))
            elif char.isAnchor() and char.anchorHref():_hyperlink(paragraph,fragment.text(),char.anchorHref())
            else:
                run=paragraph.add_run(fragment.text());run.bold=char.fontWeight()>=700;run.italic=char.fontItalic();run.underline=char.fontUnderline();run.font.name=char.font().family() or 'Calibri';run.font.size=Pt(char.fontPointSize() or 12)
                run.font.strike=char.fontStrikeOut();run.font.superscript=char.verticalAlignment()==QTextCharFormat.AlignSuperScript;run.font.subscript=char.verticalAlignment()==QTextCharFormat.AlignSubScript
                if char.background().style()!=Qt.NoBrush:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),char.background().color().name()[1:]);run._r.get_or_add_rPr().append(shade)
                if char.foreground().style()!=Qt.NoBrush:run.font.color.rgb=RGBColor.from_string(char.foreground().color().name()[1:])
        iterator+=1

def write_rich_docx(pages,path):
    original=Path(pages.imported_docx) if pages.imported_docx else None
    if original and path.resolve()==original.resolve() and pages.path is None:raise ValueError('Save an imported Word document to a new file.')
    doc=Document(original) if original else Document()
    settings=getattr(pages,'page_settings',None)
    if settings:
        from docx.shared import Mm
        from docx.enum.section import WD_ORIENT
        width,height=(210,297) if settings['paper']=='A4' else (215.9,279.4)
        for section in doc.sections:
            section.orientation=WD_ORIENT.LANDSCAPE if settings['landscape'] else WD_ORIENT.PORTRAIT
            section.page_width=Mm(height if settings['landscape'] else width);section.page_height=Mm(width if settings['landscape'] else height)
            section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Mm(settings['margin'])
    preserve_body=bool(original and pages.editor.toHtml()==getattr(pages,'original_body_html',None))
    if original and not preserve_body:
        # Existing unsupported fields/objects must not disappear during rich editing.
        bad=[el for el in doc.element.body.iter() if el.tag.rsplit('}',1)[-1] in {'fldChar','instrText','object','sdt','chart'}]
        if bad:raise ValueError('This Word file contains fields or advanced objects. Save an HTML copy for rich editing; the source remains intact.')
        for child in list(doc.element.body):
            if child.tag!=qn('w:sectPr'):doc.element.body.remove(child)
    iterator=pages.editor.document().rootFrame().begin()
    while not preserve_body and not iterator.atEnd():
        frame=iterator.currentFrame();block=iterator.currentBlock()
        if isinstance(frame,QTextTable):
            table=doc.add_table(rows=frame.rows(),cols=frame.columns());table.style='Table Grid'
            for row in range(frame.rows()):
                for col in range(frame.columns()):
                    cell=frame.cellAt(row,col);current=cell.firstCursorPosition().block();end=cell.lastCursorPosition().position();p=table.cell(row,col).paragraphs[0]
                    while current.isValid() and current.position()<=end:
                        _block(p,current,pages.editor.document());current=current.next()
                        if current.position()<=end:p=table.cell(row,col).add_paragraph()
        elif block.isValid():_block(doc.add_paragraph(style='List Bullet' if block.textList() else None),block,pages.editor.document())
        iterator+=1
    if getattr(pages,'header_footer',None) and getattr(pages,'header_footer_changed',False):
        from office_headers import apply_headers
        apply_headers(doc,pages.header_footer)
    import tempfile
    with tempfile.NamedTemporaryFile(dir=path.parent,suffix='.docx',delete=False) as stream:temporary=Path(stream.name)
    try:doc.save(temporary);temporary.replace(path)
    finally:temporary.unlink(missing_ok=True)
    pages.imported_docx=path;pages.path=path;pages.original_paragraphs=pages.editor.toPlainText().split('\n');pages.editor.document().setModified(False)
    pages.original_body_html=pages.editor.toHtml()
    pages.header_footer_changed=False

def docx_html(doc):
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    def paragraph_html(paragraph):
        parts=[]
        for item in paragraph.iter_inner_content():
            runs=[item] if hasattr(item,'font') else item.runs
            for run in runs:
                font=run.font;style=''
                if font.name:style+='font-family:'+html.escape(font.name,quote=True)+';'
                if font.size:style+=f'font-size:{font.size.pt}pt;'
                if font.bold:style+='font-weight:bold;'
                if font.italic:style+='font-style:italic;'
                decorations=('underline ' if font.underline else '')+('line-through' if font.strike else '')
                if decorations:style+='text-decoration:'+decorations.strip()+';'
                if font.superscript:style+='vertical-align:super;'
                if font.subscript:style+='vertical-align:sub;'
                if font.color.rgb:style+='color:#'+str(font.color.rgb)+';'
                shades=run._r.xpath('./w:rPr/w:shd')
                if shades:
                    color=shades[0].get(qn('w:fill'),'')
                    if len(color)==6 and all(c in '0123456789abcdefABCDEF' for c in color):style+='background-color:#'+color+';'
                text=html.escape(run.text).replace('\n','<br>')
                if hasattr(item,'url'):text='<a href="'+html.escape(item.url,quote=True)+'">'+text+'</a>'
                parts.append('<span style="'+style+'">'+text+'</span>')
                for blip in run._r.xpath('.//a:blip'):
                    rid=blip.get(qn('r:embed'))
                    if rid and rid in doc.part.related_parts:
                        image=doc.part.related_parts[rid]
                        if image.content_type.startswith('image/'):
                            parts.append('<img width="420" src="data:'+image.content_type+';base64,'+base64.b64encode(image.blob).decode()+'">')
        name=paragraph.style.name if paragraph.style else '';tag='h'+name[-1] if name in ('Heading 1','Heading 2','Heading 3') else 'p'
        align={1:'center',2:'right',3:'justify'}.get(paragraph.alignment,'left');spacing=paragraph.paragraph_format.line_spacing
        css='text-align:'+align+';'
        if paragraph.paragraph_format.page_break_before:css+='page-break-before:always;'
        if isinstance(spacing,float):css+=f'line-height:{spacing*100:g}%;'
        content=''.join(parts) or '<br>'
        if name.startswith('List '):return '<'+('ol' if name=='List Number' else 'ul')+'><li>'+content+'</li></'+('ol' if name=='List Number' else 'ul')+'>'
        return '<'+tag+' style="'+css+'">'+content+'</'+tag+'>'
    result=[]
    for node in doc.element.body:
        if node.tag==qn('w:p'):result.append(paragraph_html(Paragraph(node,doc)))
        elif node.tag==qn('w:tbl'):
            table=Table(node,doc);result.append('<table border="1" cellspacing="0" cellpadding="5">')
            for row in table.rows:result.append('<tr>'+''.join('<td>'+''.join(paragraph_html(p) for p in cell.paragraphs)+'</td>' for cell in row.cells)+'</tr>')
            result.append('</table>')
    return ''.join(result)


def play_linked_video(pages):
    links=[];block=pages.editor.document().begin()
    while block.isValid():
        it=block.begin()
        while not it.atEnd():
            fragment=it.fragment()
            if fragment.isValid():
                url=QUrl(fragment.charFormat().anchorHref())
                if url.isLocalFile() and Path(url.toLocalFile()).suffix.lower() in {'.mp4','.mov','.m4v','.webm','.avi'}:links.append(url.toLocalFile())
            it+=1
        block=block.next()
    links=list(dict.fromkeys(links))
    if not links:raise ValueError('This document has no local video links.')
    selected,ok=QInputDialog.getItem(pages,'Play linked video','Video',links,0,False)
    if ok:
        from watchit import WatchIT
        player=WatchIT();player.add_media([selected]);player.show()
        if not hasattr(pages,'video_players'):pages.video_players=[]
        pages.video_players.append(player)
