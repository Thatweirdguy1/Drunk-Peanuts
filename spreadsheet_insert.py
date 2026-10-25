"""Append native drawing parts without rebuilding an imported XLSX package."""
import copy
import io
import posixpath
import uuid
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import BarChart,LineChart,PieChart,ScatterChart,Reference,Series
from openpyxl.drawing.image import Image
from office_surgery import _active_sheet,S,R,P

CT='http://schemas.openxmlformats.org/package/2006/content-types'

def append_drawings(source,target,objects,rows,sheet_name='Sheet'):
    book=Workbook();sheet=book.active;sheet.title=sheet_name
    for row in rows:sheet.append(row)
    for obj in objects:
        if obj['type']=='image':
            image=Image(io.BytesIO(obj['blob']));image.width=obj.get('width',360);image.height=obj.get('height',220);sheet.add_image(image,obj['anchor'])
        elif obj['type']=='video':continue
        else:
            r1,c1,r2,c2=obj['range'];kind=obj['kind']
            if kind in ('Column','Bar'):chart=BarChart();chart.type='bar' if kind=='Bar' else 'col'
            elif kind=='Line':chart=LineChart()
            elif kind=='Pie':chart=PieChart()
            else:chart=ScatterChart()
            if kind=='Scatter':
                for col in range(c1+1,c2+1):chart.series.append(Series(Reference(sheet,min_col=col,min_row=r1+1,max_row=r2),Reference(sheet,min_col=c1,min_row=r1+1,max_row=r2),title=str(sheet.cell(r1,col).value)))
            else:
                chart.add_data(Reference(sheet,min_col=c1+1,max_col=c2,min_row=r1,max_row=r2),titles_from_data=True);chart.set_categories(Reference(sheet,min_col=c1,min_row=r1+1,max_row=r2))
            chart.title=obj.get('title','Chart');chart.width=18;chart.height=10;sheet.add_chart(chart,obj['anchor'])
    stream=io.BytesIO();book.save(stream);book.close();stream.seek(0)
    with ZipFile(source) as old,ZipFile(stream) as new:
        parts={n:old.read(n) for n in old.namelist()};sheetpath=_active_sheet(old)
        sheetxml=etree.fromstring(parts[sheetpath]);sheetrel=posixpath.join(posixpath.dirname(sheetpath),'_rels',posixpath.basename(sheetpath)+'.rels')
        rels=etree.fromstring(parts[sheetrel]) if sheetrel in parts else etree.Element('{'+P+'}Relationships',nsmap={None:P})
        tag=lambda n:'{'+S+'}'+n
        if any(obj['type']!='video' for obj in objects):
            token='fc'+uuid.uuid4().hex[:12]+'_'
            selected=[n for n in new.namelist() if n.startswith(('xl/drawings/','xl/charts/','xl/media/'))]
            mapping={n:posixpath.join(posixpath.dirname(n),token+posixpath.basename(n)) for n in selected if '/_rels/' not in n}
            for name in selected:
                if '/_rels/' in name:
                    base=posixpath.join(posixpath.dirname(posixpath.dirname(name)),posixpath.basename(name)[:-5]);mapped=mapping[base];mapping[name]=posixpath.join(posixpath.dirname(mapped),'_rels',posixpath.basename(mapped)+'.rels')
            for name in selected:
                data=new.read(name)
                if name.endswith('.rels'):
                    root=etree.fromstring(data);parent=posixpath.dirname(posixpath.dirname(name))
                    for rel in root:
                        original=rel.get('Target','');resolved=original.lstrip('/') if original.startswith('/') else posixpath.normpath(posixpath.join(parent,original))
                        if resolved in mapping:rel.set('Target','/'+mapping[resolved])
                    data=etree.tostring(root)
                parts[mapping[name]]=data
            existing=sheetxml.find(tag('drawing'))
            incoming='xl/drawings/drawing1.xml';incomingrels='xl/drawings/_rels/drawing1.xml.rels'
            if existing is not None:
                rid=existing.get('{'+R+'}id');link=next(r for r in rels if r.get('Id')==rid);path=link.get('Target');drawingpath=path.lstrip('/') if path.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(sheetpath),path))
                drawing=etree.fromstring(parts[drawingpath]);drawingrel=posixpath.join(posixpath.dirname(drawingpath),'_rels',posixpath.basename(drawingpath)+'.rels')
                oldrels=etree.fromstring(parts[drawingrel]) if drawingrel in parts else etree.Element('{'+P+'}Relationships',nsmap={None:P})
                added=etree.fromstring(parts[mapping[incoming]]);addedrels=etree.fromstring(parts[mapping[incomingrels]]) if incomingrels in mapping else []
                used={r.get('Id') for r in oldrels};rename={}
                for rel in addedrels:
                    index=1
                    while 'rIdFC'+str(index) in used:index+=1
                    new_id='rIdFC'+str(index);used.add(new_id);rename[rel.get('Id')]=new_id;rel.set('Id',new_id);oldrels.append(rel)
                ids=[int(el.get('id')) for el in drawing.iter() if etree.QName(el).localname=='cNvPr' and el.get('id','').isdigit()];next_id=max(ids,default=0)+1
                for anchor in added:
                    for el in anchor.iter():
                        for attr,value in list(el.attrib.items()):
                            if attr in ('{'+R+'}id','{'+R+'}embed','{'+R+'}link') and value in rename:el.set(attr,rename[value])
                        if etree.QName(el).localname=='cNvPr':el.set('id',str(next_id));next_id+=1
                    drawing.append(anchor)
                parts[drawingpath]=etree.tostring(drawing);parts[drawingrel]=etree.tostring(oldrels)
                parts.pop(mapping[incoming],None);parts.pop(mapping.get(incomingrels,''),None)
            else:
                index=1;used={r.get('Id') for r in rels}
                while 'rIdFC'+str(index) in used:index+=1
                rid='rIdFC'+str(index);rel=etree.SubElement(rels,'{'+P+'}Relationship');rel.set('Id',rid);rel.set('Type',R+'/drawing');rel.set('Target','/'+mapping[incoming]);node=etree.Element(tag('drawing'));node.set('{'+R+'}id',rid)
                following={'legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst'}
                index=next((i for i,n in enumerate(sheetxml) if etree.QName(n).localname in following),len(sheetxml));sheetxml.insert(index,node)
            types=etree.fromstring(parts['[Content_Types].xml']);newtypes=etree.fromstring(new.read('[Content_Types].xml'));known={n.get('PartName') or n.get('Extension') for n in types}
            for node in newtypes:
                name=node.get('PartName','').lstrip('/')
                if name in mapping and mapping[name] in parts:node.set('PartName','/'+mapping[name])
                elif node.get('Extension') and node.get('Extension') not in known:pass
                else:continue
                key=node.get('PartName') or node.get('Extension')
                if key not in known:types.append(node);known.add(key)
            parts['[Content_Types].xml']=etree.tostring(types)
        for obj in objects:
            if obj['type']!='video':continue
            links=sheetxml.find(tag('hyperlinks'))
            if links is None:
                links=etree.Element(tag('hyperlinks'));sheetxml.insert(next((i for i,n in enumerate(sheetxml) if etree.QName(n).localname in {'printOptions','pageMargins','drawing','extLst'}),len(sheetxml)),links)
            rid='rIdMedia'+uuid.uuid4().hex[:10];rel=etree.SubElement(rels,'{'+P+'}Relationship');rel.set('Id',rid);rel.set('Type',R+'/hyperlink');rel.set('Target',obj['url']);rel.set('TargetMode','External');node=etree.SubElement(links,tag('hyperlink'));node.set('ref',obj['anchor']);node.set('{'+R+'}id',rid)
        parts[sheetpath]=etree.tostring(sheetxml);parts[sheetrel]=etree.tostring(rels)
        with ZipFile(target,'w',ZIP_DEFLATED) as output:
            for name,data in parts.items():output.writestr(name,data)
