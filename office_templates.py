"""Bundled editable starters, copied before use so originals stay intact."""
from pathlib import Path
import shutil

ASSETS = Path(__file__).resolve().parent / 'assets'
TEMPLATES = {
    'Pages': ('Report.docx', 'Letter.docx', 'Meeting notes.docx'),
    'Tables': ('Budget.xlsx', 'Task tracker.xlsx'),
    'Dazzle': ('Presentation.pptx', 'Project update.pptx', 'Pitch deck.pptx', 'Research presentation.pptx', 'Lesson plan.pptx', 'Portfolio.pptx'),
    'Sketch': ('Blank canvas.png',),
    'Organize': ('Contacts.sqlite',),
    'Formulaic': ('Equations.formulaic',),
    "Writer's Drafts": ('Screenplay.fountain',),
}

SLIDE_COLORS = dict(background='#0b0b0b', title='#ffdd22', body='#f2f2ed')


def style_slide(slide, colors=None):
    """Give editable title/body placeholders the selected slide colors."""
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    colors = colors or SLIDE_COLORS
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(colors['background'][1:])
    for shape, role, size in ((slide.shapes.title, 'title', 32), (slide.placeholders[1], 'body', 20)):
        for paragraph in shape.text_frame.paragraphs:
            paragraph.font.name = 'Calibri'
            paragraph.font.size = Pt(size)
            paragraph.font.bold = role == 'title'
            paragraph.font.color.rgb = RGBColor.from_string(colors[role][1:])


def create_document(workspace, name, destination):
    if name not in TEMPLATES.get(workspace, ()):
        raise ValueError('Choose a bundled template')
    destination = Path(destination)
    source = ASSETS / 'templates' / name
    if destination.suffix.lower() != source.suffix.lower():
        raise ValueError(f'This template needs a {source.suffix} filename')
    with source.open('rb') as original, destination.open('xb') as copy:
        shutil.copyfileobj(original, copy)
    return destination
