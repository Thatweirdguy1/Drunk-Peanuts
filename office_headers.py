"""Native Word header/footer stories and page fields, plus local PDF overlays."""
import io,tempfile
from pathlib import Path
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QDialog,QFormLayout,QPlainTextEdit,QComboBox,QCheckBox,QSpinBox,QDialogButtonBox,QLabel,QMessageBox
from office_editing import command


def validate(settings):
    if not all(isinstance(settings.get(k),str) and len(settings[k])<=2000 for k in ('header','footer')):raise ValueError('Header and footer text must be at most 2,000 characters each.')
    if settings.get('numbers') not in ('None','Header','Footer') or settings.get('align') not in ('Left','Center','Right'):raise ValueError('Choose a supported position and alignment.')
    if not isinstance(settings.get('start'),int) or not 1<=settings['start']<=9999:raise ValueError('Start numbering between 1 and 9999.')


def set_headers(pages,settings):
    validate(settings);pages.header_footer=dict(settings);pages.rich_changed=True;pages.editor.document().setModified(True)
    for location,preview in [('Header',pages.header_preview),('Footer',pages.footer_preview)]:
        text=settings[location.lower()]
        if settings['numbers']==location:text+=(' · ' if text else '')+f'Page {settings["start"]} of …'
        preview.setText(text);preview.setAlignment({'Left':Qt.AlignLeft,'Center':Qt.AlignHCenter,'Right':Qt.AlignRight}[settings['align']]);preview.setVisible(bool(text) and not settings['skip_first'])
        preview.setToolTip('First-page preview. Final page counts appear in DOCX/PDF output.')
    pages.statusBar().showMessage('Header/footer and page numbers will appear in DOCX and PDF output.')


def edit_headers(pages):
    current=pages.header_footer or dict(header='',footer='',numbers='Footer',align='Center',start=1,skip_first=False)
    if pages.header_footer is None and pages.imported_docx:
        from docx import Document
        section=Document(pages.imported_docx).sections[0]
        current.update(header='\n'.join(p.text for p in section.header.paragraphs),footer='\n'.join(p.text for p in section.footer.paragraphs))
    dialog=QDialog(pages);dialog.setWindowTitle('Header, footer and page numbers');form=QFormLayout(dialog)
    header=QPlainTextEdit(current['header']);footer=QPlainTextEdit(current['footer']);header.setMaximumHeight(90);footer.setMaximumHeight(90);header.setAccessibleName('Header text');footer.setAccessibleName('Footer text')
    numbers=QComboBox();numbers.addItems(['None','Header','Footer']);numbers.setCurrentText(current['numbers']);align=QComboBox();align.addItems(['Left','Center','Right']);align.setCurrentText(current['align']);start=QSpinBox();start.setRange(1,9999);start.setValue(current['start']);skip=QCheckBox('Hide header/footer on the first page');skip.setChecked(current['skip_first'])
    form.addRow('Header',header);form.addRow('Footer',footer);form.addRow('Page numbers',numbers);form.addRow('Alignment',align);form.addRow('Start at',start);form.addRow(skip)
    note=QLabel('Applies to every section. Replaces existing header/footer stories in the saved copy.\nDOCX uses live page fields; PDF uses the exported page count.');note.setWordWrap(True);form.addRow(note)
    buttons=QDialogButtonBox(QDialogButtonBox.Ok|QDialogButtonBox.Cancel);form.addRow(buttons)
    def accept():
        settings=dict(header=header.toPlainText(),footer=footer.toPlainText(),numbers=numbers.currentText(),align=align.currentText(),start=start.value(),skip_first=skip.isChecked())
        try:set_headers(pages,settings)
        except ValueError as error:QMessageBox.warning(dialog,'Header/footer',str(error));return
        dialog.accept()
    buttons.accepted.connect(accept);buttons.rejected.connect(dialog.reject);dialog.resize(520,430);dialog.exec()


def apply_headers(document,settings):
    validate(settings)
    def field(paragraph,instruction):
        run=paragraph.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');code=OxmlElement('w:instrText');code.set(qn('xml:space'),'preserve');code.text=' '+instruction+' ';end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');run._r.extend([begin,code,end])
    for index,section in enumerate(document.sections):
        section.different_first_page_header_footer=bool(settings['skip_first'])
        section.even_page_header.is_linked_to_previous=False;section.even_page_footer.is_linked_to_previous=False
        for location,text in [('Header',settings['header']),('Footer',settings['footer'])]:
            for story in ([section.header,section.even_page_header] if location=='Header' else [section.footer,section.even_page_footer]):
                story.is_linked_to_previous=False
                for child in list(story._element):story._element.remove(child)
                paragraph=story.add_paragraph(text);paragraph.alignment={'Left':WD_ALIGN_PARAGRAPH.LEFT,'Center':WD_ALIGN_PARAGRAPH.CENTER,'Right':WD_ALIGN_PARAGRAPH.RIGHT}[settings['align']]
                if settings['numbers']==location:
                    if text:paragraph.add_run(' · ')
                    paragraph.add_run('Page ');field(paragraph,'PAGE');paragraph.add_run(' of ');field(paragraph,'NUMPAGES')
            first=section.first_page_header if location=='Header' else section.first_page_footer;first.is_linked_to_previous=False
            for child in list(first._element):first._element.remove(child)
            first.add_paragraph('')
        existing=section._sectPr.find(qn('w:pgNumType'))
        if existing is not None:section._sectPr.remove(existing)
        if index==0:
            node=OxmlElement('w:pgNumType');node.set(qn('w:start'),str(settings['start']));section._sectPr.append(node)
    update=document.settings.element.find(qn('w:updateFields'))
    if update is None:update=OxmlElement('w:updateFields');document.settings.element.append(update)
    update.set(qn('w:val'),'true')


def overlay_pdf(path,settings):
    validate(settings)
    from pypdf import PdfReader,PdfWriter
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import simpleSplit
    reader=PdfReader(str(path));writer=PdfWriter();total=len(reader.pages)
    overlays=[]
    for index,page in enumerate(reader.pages):
        page=writer.add_page(page)
        if not (index==0 and settings['skip_first']):
            width,height=float(page.mediabox.width),float(page.mediabox.height);stream=io.BytesIO();drawing=canvas.Canvas(stream,pagesize=(width,height));drawing.setFont('Helvetica',9)
            for location,text,y in [('Header',settings['header'],height-25),('Footer',settings['footer'],25)]:
                if settings['numbers']==location:text+=((' · ' if text else '')+f'Page {settings["start"]+index} of {total}')
                lines=[wrapped for line in text.splitlines() for wrapped in simpleSplit(line,'Helvetica',9,width-80)]
                if location=='Footer':y+=11*max(0,len(lines)-1)
                for wrapped in lines:
                    if settings['align']=='Left':drawing.drawString(40,y,wrapped)
                    elif settings['align']=='Right':drawing.drawRightString(width-40,y,wrapped)
                    else:drawing.drawCentredString(width/2,y,wrapped)
                    y-=11
            drawing.save();stream.seek(0);overlay=PdfReader(stream);overlays.append(overlay);page.merge_page(overlay.pages[0])
    with tempfile.NamedTemporaryFile(dir=Path(path).parent,suffix='.pdf',delete=False) as stream:temporary=Path(stream.name)
    try:
        with temporary.open('wb') as output:writer.write(output)
        temporary.replace(path)
    finally:writer.close();temporary.unlink(missing_ok=True)


def setup_header_tools(pages):
    pages.header_footer=None;pages.original_body_html=None
    for label in ('Header','Footer','Page numbers'):command(pages,'Insert',label,lambda:edit_headers(pages),'title')
