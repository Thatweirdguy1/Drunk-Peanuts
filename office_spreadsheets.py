"""Spreadsheet insertion UI; real charts and images are stored as XLSX drawings."""
import io
import tempfile
from pathlib import Path
from PySide6.QtCore import Qt
from PySide6.QtGui import QImage,QPixmap
from PySide6.QtWidgets import QDockWidget,QWidget,QVBoxLayout,QListWidget,QLabel,QFileDialog,QInputDialog,QTableWidgetItem,QPushButton
from openpyxl.utils import get_column_letter
from office_insert import image_bytes,chart_png
from spreadsheet_insert import append_drawings

def setup_tables(editor):
    editor.inserted_objects=[]
    dock=QDockWidget('Charts and images',editor);dock.setObjectName('sheetObjects');panel=QWidget();layout=QVBoxLayout(panel)
    editor.object_list=QListWidget();editor.object_list.setMaximumHeight(140);layout.addWidget(editor.object_list)
    editor.object_preview=QLabel('Select an inserted object');editor.object_preview.setAlignment(Qt.AlignCenter);editor.object_preview.setWordWrap(True);layout.addWidget(editor.object_preview);play=QPushButton('Play selected video link');play.clicked.connect(lambda:editor.report(lambda:play_link(editor)));layout.addWidget(play);dock.setWidget(panel);editor.addDockWidget(Qt.RightDockWidgetArea,dock);dock.setMinimumWidth(230)
    editor.object_list.currentRowChanged.connect(lambda _:preview(editor))
    menu=editor.menuBar().addMenu('Insert')
    for label,callback in [('Image',lambda:insert_image(editor)),('Chart from selected cells',lambda:insert_chart(editor)),('Local video link',lambda:insert_video(editor))]:menu.addAction(label,lambda checked=False,cb=callback:editor.report(cb))
    editor.grid.itemChanged.connect(lambda _:preview(editor))

def anchor(editor):return f'{get_column_letter(max(0,editor.grid.currentColumn())+1)}{max(0,editor.grid.currentRow())+1}'

def add(editor,obj):
    editor.inserted_objects.append(obj);editor.object_list.addItem(obj['type'].title()+' · '+obj['anchor']);editor.object_list.setCurrentRow(len(editor.inserted_objects)-1);editor.dirty=True

def insert_image(editor,path=None):
    if path is None:
        name,_=QFileDialog.getOpenFileName(editor,'Insert image','','Images (*.png *.jpg *.jpeg *.webp *.bmp)');path=Path(name) if name else None
    if not path:return
    blob=image_bytes(path);image=QImage.fromData(blob);add(editor,dict(type='image',blob=blob,anchor=anchor(editor),width=360,height=360*image.height()/max(1,image.width())))

def chart_values(editor,bounds):
    r1,c1,r2,c2=bounds
    categories=[editor.grid.item(r-1,c1-1).text() if editor.grid.item(r-1,c1-1) else str(r) for r in range(r1+1,r2+1)]
    series=[]
    for col in range(c1+1,c2+1):
        name=editor.grid.item(r1-1,col-1);values=[]
        for row in range(r1+1,r2+1):
            item=editor.grid.item(row-1,col-1);values.append(float(item.text()) if item and item.text() else 0)
        series.append({'name':name.text() if name else get_column_letter(col),'values':values})
    return dict(categories=categories,series=series)

def insert_chart(editor,kind=None,bounds=None):
    if bounds is None:
        ranges=editor.grid.selectedRanges()
        if len(ranges)!=1:raise ValueError('Select one rectangular data range, with headers and categories.')
        selected=ranges[0];bounds=(selected.topRow()+1,selected.leftColumn()+1,selected.bottomRow()+1,selected.rightColumn()+1)
    r1,c1,r2,c2=bounds
    if r2<=r1 or c2<=c1 or r2-r1>500 or c2-c1>8:raise ValueError('Select a header row, category column and up to 500 rows / 8 numeric series.')
    if kind is None:
        kind,ok=QInputDialog.getItem(editor,'Insert chart','Type',['Column','Bar','Line','Pie','Scatter'],0,False)
        if not ok:return
    data=chart_values(editor,bounds);chart_png(data,kind)
    add(editor,dict(type='chart',kind=kind,range=bounds,anchor=anchor(editor),title='Chart'))

def insert_video(editor,path=None):
    if path is None:
        name,_=QFileDialog.getOpenFileName(editor,'Insert local video link','','Video (*.mp4 *.mov *.webm *.avi)');path=Path(name) if name else None
    if not path:return
    path=Path(path).resolve()
    if not path.is_file():raise ValueError('Video file was not found.')
    target=anchor(editor);editor.grid.setItem(editor.grid.currentRow(),editor.grid.currentColumn(),QTableWidgetItem('Video: '+path.name));add(editor,dict(type='video',anchor=target,url=path.as_uri()))

def preview(editor):
    index=editor.object_list.currentRow()
    if not 0<=index<len(editor.inserted_objects):return
    obj=editor.inserted_objects[index]
    try:
        if obj['type']=='image':blob=obj['blob']
        elif obj['type']=='chart':blob=chart_png(chart_values(editor,obj['range']),obj['kind'])
        else:editor.object_preview.setText('Local video link at '+obj['anchor']+'\n'+obj['url']);return
        editor.object_preview.setPixmap(QPixmap.fromImage(QImage.fromData(blob)).scaled(300,240,Qt.KeepAspectRatio,Qt.SmoothTransformation))
    except ValueError:editor.object_preview.setText('Chart data needs numeric values.')

def save_objects(editor,path,rows):
    pending=[obj for obj in editor.inserted_objects if not obj.get('saved')]
    if not pending:return
    from openpyxl import load_workbook
    book=load_workbook(path,read_only=True);name=book.active.title;book.close()
    with tempfile.NamedTemporaryFile(dir=path.parent,suffix='.xlsx',delete=False) as stream:temporary=Path(stream.name)
    try:append_drawings(path,temporary,pending,rows,name);temporary.replace(path)
    finally:temporary.unlink(missing_ok=True)
    for obj in pending:obj['saved']=True

def load_objects(editor,path):
    """Read drawing metadata directly; never resave an imported workbook through openpyxl."""
    from zipfile import ZipFile
    from lxml import etree
    import posixpath
    from office_surgery import _active_sheet,R,P,S
    from openpyxl.utils.cell import range_boundaries
    def resolve(part,target):return target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(part),target))
    def relationships(archive,part):
        relpath=posixpath.join(posixpath.dirname(part),'_rels',posixpath.basename(part)+'.rels')
        if relpath not in archive.namelist():return {}
        return {n.get('Id'):n.get('Target') for n in etree.fromstring(archive.read(relpath))}
    ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart','a':'http://schemas.openxmlformats.org/drawingml/2006/main','xdr':'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'}
    with ZipFile(path) as archive:
        sheetpath=_active_sheet(archive);sheet=etree.fromstring(archive.read(sheetpath));rels=relationships(archive,sheetpath);drawing=sheet.find('{'+S+'}drawing')
        if drawing is not None:
            drawingpath=resolve(sheetpath,rels[drawing.get('{'+R+'}id')]);links=relationships(archive,drawingpath)
            for obj in etree.fromstring(archive.read(drawingpath)):
                row=obj.find('xdr:from/xdr:row',ns);col=obj.find('xdr:from/xdr:col',ns);position=f'{get_column_letter(int(col.text)+1)}{int(row.text)+1}' if row is not None and col is not None else 'A1'
                blip=obj.find('.//a:blip',ns);chart=obj.find('.//c:chart',ns)
                if blip is not None:
                    part=resolve(drawingpath,links[blip.get('{'+R+'}embed')]);add(editor,dict(type='image',blob=archive.read(part),anchor=position,saved=True))
                elif chart is not None:
                    root=etree.fromstring(archive.read(resolve(drawingpath,links[chart.get('{'+R+'}id')])));cat=root.find('.//c:cat/*/c:f',ns);values=root.findall('.//c:val/*/c:f',ns)
                    if cat is None:cat=root.find('.//c:xVal/*/c:f',ns);values=root.findall('.//c:yVal/*/c:f',ns)
                    if cat is None or not values:continue
                    c1,r1,cx,r2=range_boundaries(cat.text.rsplit('!',1)[-1]);c2=max(range_boundaries(v.text.rsplit('!',1)[-1])[2] for v in values)
                    kind='Line' if root.find('.//c:lineChart',ns) is not None else 'Pie' if root.find('.//c:pieChart',ns) is not None else 'Scatter' if root.find('.//c:scatterChart',ns) is not None else 'Column'
                    direction=root.find('.//c:barDir',ns)
                    if direction is not None and direction.get('val')=='bar':kind='Bar'
                    add(editor,dict(type='chart',kind=kind,range=(max(1,r1-1),c1,r2,c2),anchor=position,saved=True))
        for link in sheet.findall('{'+S+'}hyperlinks/{'+S+'}hyperlink'):
            url=rels.get(link.get('{'+R+'}id'),'')
            if url.startswith('file:'):add(editor,dict(type='video',anchor=link.get('ref'),url=url,saved=True))
    editor.dirty=False


def play_link(editor):
    from PySide6.QtCore import QUrl
    index=editor.object_list.currentRow()
    if not 0<=index<len(editor.inserted_objects) or editor.inserted_objects[index]['type']!='video':raise ValueError('Select a video link in the object list.')
    url=QUrl(editor.inserted_objects[index]['url'])
    if not url.isLocalFile():raise ValueError('Only local video links can be played here.')
    from watchit import WatchIT
    player=WatchIT();player.add_media([url.toLocalFile()]);player.show()
    if not hasattr(editor,'video_players'):editor.video_players=[]
    editor.video_players.append(player)
