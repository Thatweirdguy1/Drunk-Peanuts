import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import io,tempfile,unittest
from pathlib import Path
from PIL import Image,ImageDraw,ImageChops,ImageStat
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from docx import Document
from pypdf import PdfReader
import pypdfium2 as pdfium
from converter import convert


class PdfDocxPreservationTests(unittest.TestCase):
    def test_default_preserves_graphical_header_and_explicit_text_mode_is_separate(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);pdf=root/'sample.pdf';docx=root/'preserved.docx'
            header=Image.new('RGB',(500,100),'#371652');ImageDraw.Draw(header).text((30,30),'GRAPHIC LETTERHEAD',fill='yellow');buffer=io.BytesIO();header.save(buffer,format='PNG');buffer.seek(0)
            page=canvas.Canvas(str(pdf),pagesize=(600,800));page.drawImage(ImageReader(buffer),40,670,width=500,height=100);page.drawString(45,600,'Editable sample body');page.save();before=pdf.read_bytes()
            convert(pdf,docx);result=Document(docx)
            self.assertIn('Editable sample body',' '.join(p.text for p in result.paragraphs))
            from zipfile import ZipFile
            with ZipFile(docx) as archive:
                self.assertIn(b'wp:anchor',archive.read('word/document.xml'))
                data=archive.read(next(n for n in archive.namelist() if n.startswith('word/media/')))
            actual=Image.open(io.BytesIO(data)).convert('RGB');self.assertEqual(actual.size,header.size)
            self.assertLess(max(ImageStat.Stat(ImageChops.difference(actual,header)).mean),.01)
            preserved=root/'visual.docx';convert(pdf,preserved,'visual');self.assertEqual(len(Document(preserved).inline_shapes),1)
            with pdfium.PdfDocument(str(pdf)) as original:
                rendered=original[0].render(scale=1.5).to_pil().convert('RGB')
            from zipfile import ZipFile
            with ZipFile(preserved) as archive:
                data=archive.read(next(n for n in archive.namelist() if n.startswith('word/media/')))
            actual=Image.open(io.BytesIO(data)).convert('RGB');self.assertEqual(actual.size,rendered.size);self.assertLess(max(ImageStat.Stat(ImageChops.difference(actual,rendered)).mean),.01)
            text=root/'text.docx';convert(pdf,text,'text');plain=Document(text);self.assertEqual(len(plain.inline_shapes),0);self.assertIn('Editable sample body',' '.join(p.text for p in plain.paragraphs));self.assertEqual(pdf.read_bytes(),before)


if __name__=='__main__':unittest.main()
