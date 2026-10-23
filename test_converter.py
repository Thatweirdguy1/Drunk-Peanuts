"""Small end-to-end checks using only temporary files."""

import os
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
import wave

from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from PySide6.QtCore import QMimeData, QPointF, Qt, QUrl
from PySide6.QtGui import QDropEvent
from PySide6.QtWidgets import QApplication

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from converter import available_formats, compress, convert, reduce_size
from office_pdf import available_engines
from app import Window, Worker, numbered_target
from office_suite import calculate_cell
from converter import edit_image, extract_zip, ocr_to_text, pdf_merge, pdf_reorder, pdf_rotate, pdf_split, trim_media, zip_entries


class ConversionTest(unittest.TestCase):
    def test_office_formats_are_native(self):
        self.assertIn("pdf", available_formats(Path("slides.pptx")))
        self.assertIn("odt", available_formats(Path("letter.docx")))
        self.assertEqual(available_formats(Path("old.doc")), ["docx", "pdf"] if available_engines(Path("old.doc")) else [])

    def test_image_pdf_and_back(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "sample.png"
            Image.new("RGB", (32, 24), "#0d9488").save(source)
            pdf = root / "sample.pdf"
            convert(source, pdf)
            self.assertEqual(len(PdfReader(str(pdf)).pages), 1)
            output = root / "rendered.png"
            convert(pdf, output)
            with Image.open(output) as image:
                self.assertEqual(image.size, (64, 48))
            with self.assertRaises(FileExistsError):
                convert(source, pdf)

    def test_pdf_visual_docx_keeps_page_images(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source.txt"
            source.write_text("First page\n" * 95, encoding="utf-8")
            pdf = root / "source.pdf"
            convert(source, pdf)
            visual = root / "visual.docx"
            convert(pdf, visual, "visual")
            self.assertEqual(len(Document(visual).inline_shapes), len(PdfReader(str(pdf)).pages))
            app = QApplication.instance() or QApplication([])
            window = Window()
            window.output_dir = root
            window._add_files([pdf])
            window.formats[pdf.resolve()] = "docx"
            window.engines[pdf.resolve()] = "visual"
            window._convert()
            self.assertTrue(window.worker.wait(10000))
            app.processEvents()
            self.assertEqual(len(Document(root / "source.docx").inline_shapes), len(PdfReader(str(pdf)).pages))
            window.close()

    def test_animated_image_keeps_frames(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "animated.gif"
            Image.new("RGB", (20, 20), "red").save(source, save_all=True,
                append_images=[Image.new("RGB", (20, 20), "blue")], duration=[100, 150])
            self.assertNotIn("jpg", available_formats(source))
            for format in ("webp", "tiff"):
                output = root / f"animated.{format}"
                convert(source, output)
                with Image.open(output) as image:
                    self.assertEqual(image.n_frames, 2)
            edited = root / "resized.gif"
            edit_image(source, edited, "resize", (10, 10))
            with Image.open(edited) as image:
                self.assertEqual((image.n_frames, image.size), (2, (10, 10)))
                self.assertEqual([image.seek(i) or image.info["duration"] for i in range(2)], [100, 150])

    def test_csv_xlsx_roundtrip(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "data.csv"
            source.write_text("name,value\nalpha,3\n", encoding="utf-8")
            spreadsheet = root / "data.xlsx"
            convert(source, spreadsheet)
            output = root / "again.csv"
            convert(spreadsheet, output)
            self.assertIn("alpha,3", output.read_text(encoding="utf-8-sig"))
            self.assertNotIn("xlsx", available_formats(spreadsheet))
            semicolon = root / "european.csv"
            semicolon.write_text('name;value\n"Smith; Jones";4\n', encoding="utf-8")
            converted = root / "european.xlsx"
            convert(semicolon, converted)
            workbook = load_workbook(converted)
            self.assertEqual([cell.value for cell in workbook.active[2]], ["Smith; Jones", "4"])
            workbook.close()

    def test_table_formula_calculation(self):
        rows = [["2", "=A1*3"], ["4", "=SUM(A1:A2)+B1"]]
        self.assertEqual(calculate_cell(rows, "B2"), 12)
        with self.assertRaises(ValueError):
            calculate_cell([["=A1+1"]], "A1")
        with self.assertRaises(ValueError):
            calculate_cell([["=__import__('os').system('echo unsafe')"]], "A1")

    def test_local_document_conversion(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "notes.txt"
            source.write_text("First line\nSecond line", encoding="utf-8")
            word = root / "notes.docx"
            convert(source, word)
            self.assertEqual([p.text for p in Document(word).paragraphs], ["First line", "Second line"])
            pdf = root / "notes.pdf"
            convert(source, pdf)
            extracted = root / "extracted.docx"
            convert(pdf, extracted)
            self.assertIn("First line", "\n".join(p.text for p in Document(extracted).paragraphs))
            self.assertIn("docx", available_formats(pdf))

        self.assertNotIn("xlsx", available_formats(Path("slides.pptx")))
        self.assertIn("rtf", available_formats(Path("letter.docx")))

    def test_native_office_roundtrips(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "note.txt"
            source.write_text("Hello native Office", encoding="utf-8")
            for fmt in ("docx", "odt", "rtf"):
                middle = root / f"note.{fmt}"
                convert(source, middle)
                output = root / f"from-{fmt}.txt"
                convert(middle, output)
                self.assertIn("Hello native Office", output.read_text(encoding="utf-8"))
            csv_source = root / "sheet.csv"
            csv_source.write_text("name,value\nitem,12\n", encoding="utf-8")
            ods = root / "sheet.ods"
            convert(csv_source, ods)
            xlsx = root / "sheet.xlsx"
            convert(ods, xlsx)
            book = load_workbook(xlsx, read_only=True)
            self.assertEqual(book.active["A2"].value, "item")
            book.close()
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = "Opening slide"
            pptx = root / "slides.pptx"
            deck.save(pptx)
            odp = root / "slides.odp"
            convert(pptx, odp)
            converted = root / "again.pptx"
            convert(odp, converted)
            self.assertIn("Opening slide", [shape.text for shape in Presentation(converted).slides[0].shapes if shape.has_text_frame])

    def test_advanced_office_conversion_refuses_loss(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            book = __import__("openpyxl").Workbook()
            book.active["A1"] = "=2+2"
            source = root / "formula.xlsx"
            book.save(source)
            with self.assertRaisesRegex(ValueError, "formulas"):
                convert(source, root / "formula.ods")
            self.assertFalse((root / "formula.ods").exists())
            doc = Document()
            image = root / "image.png"
            Image.new("RGB", (8, 8), "red").save(image)
            doc.add_picture(str(image))
            rich = root / "document.docx"
            doc.save(rich)
            with self.assertRaisesRegex(ValueError, "word/media"):
                convert(rich, root / "document.odt")
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_shape(1, 0, 0, 914400, 914400)
            pptx = root / "shapes.pptx"
            deck.save(pptx)
            with self.assertRaisesRegex(ValueError, "nontext slide shapes"):
                convert(pptx, root / "shapes.odp")

    @unittest.skipUnless("libreoffice" in available_engines(Path("sample.docx")), "LibreOffice is not installed")
    def test_installed_libreoffice_keeps_image_on_cross_format_conversion(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            picture = root / "picture.png"
            Image.new("RGB", (24, 24), "red").save(picture)
            document = Document()
            document.add_paragraph("Image and text")
            document.add_picture(str(picture))
            source = root / "rich.docx"
            document.save(source)
            target = root / "rich.odt"
            convert(source, target, "libreoffice")
            with ZipFile(target) as archive:
                self.assertTrue(any(name.startswith("Pictures/") for name in archive.namelist()))
                self.assertIn("Image and text", archive.read("content.xml").decode())
            app = QApplication.instance() or QApplication([])
            window = Window()
            window.output_dir = root
            window._add_files([source])
            window.formats[source.resolve()] = "odt"
            window.engines[source.resolve()] = "libreoffice"
            with patch("app.convert") as run_conversion:
                window._convert()
                self.assertTrue(window.worker.wait(10000))
            app.processEvents()
            self.assertEqual(run_conversion.call_args.args[2], "libreoffice")
            window.close()

    def test_zip_compression(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            first, second = root / "one.txt", root / "two.txt"
            first.write_text("hello" * 100, encoding="utf-8")
            second.write_text("world" * 100, encoding="utf-8")
            archive = root / "files.zip"
            compress([first, second], archive)
            with ZipFile(archive) as opened:
                self.assertEqual(opened.namelist(), ["one.txt", "two.txt"])
                self.assertEqual(opened.read("one.txt"), first.read_bytes())
            with self.assertRaises(FileExistsError):
                compress([first], archive)
            other = root / "other"
            other.mkdir()
            duplicate = other / "one.txt"
            duplicate.write_text("different", encoding="utf-8")
            numbered = root / "numbered.zip"
            compress([first, duplicate], numbered)
            with ZipFile(numbered) as opened:
                self.assertEqual(opened.namelist(), ["one.txt", "one (2).txt"])
                self.assertEqual(opened.read("one (2).txt"), b"different")

    def test_target_size_image(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "large.png"
            Image.effect_noise((256, 256), 100).convert("RGB").save(source)
            target = root / "large-smaller.webp"
            reduce_size(source, target, 0.02)
            self.assertLessEqual(target.stat().st_size, int(0.02 * 1024 * 1024))
            with Image.open(target) as image:
                self.assertEqual(image.format, "WEBP")
            animated = root / "large.gif"
            frames = [Image.effect_noise((256, 256), seed).convert("RGB") for seed in (90, 110)]
            frames[0].save(animated, save_all=True, append_images=frames[1:], duration=[100, 150])
            smaller = root / "animated-smaller.webp"
            reduce_size(animated, smaller, 0.08)
            with Image.open(smaller) as image:
                self.assertEqual(image.n_frames, 2)
            self.assertLessEqual(smaller.stat().st_size, int(0.08 * 1024 * 1024))

    def test_pdf_size_reduction_keeps_text_when_lossless_fits(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "large.pdf"
            document = canvas.Canvas(str(source), pageCompression=0)
            for line in range(100):
                document.drawString(30, 800 - line * 7, "SELECTABLE TEXT " * 8)
            document.save()
            output = root / "small.pdf"
            reduce_size(source, output, source.stat().st_size * 0.8 / 1024 / 1024)
            self.assertIn("SELECTABLE TEXT", PdfReader(str(output)).pages[0].extract_text())

    def test_batch_runs_two_files_concurrently(self):
        barrier = threading.Barrier(2, timeout=3)
        worker = Worker([(Path("a"), Path("a.out")), (Path("b"), Path("b.out"))])
        result = []
        worker.finished_batch.connect(lambda done, failed: result.append((done, failed)))
        with patch("app.convert", side_effect=lambda *_: barrier.wait()):
            worker.run()
        self.assertEqual(result, [(2, 0)])

    def test_shift_drop_chooses_format_and_converts(self):
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "note.txt"
            source.write_text("hello", encoding="utf-8")
            window = Window()

            class Menu:
                def __init__(self, *_):
                    self.actions = []

                def addAction(self, label):
                    self.actions.append(label)
                    return label

                def exec(self, _):
                    return "DOCX"

            with patch("app.QMenu", Menu), patch.object(window, "_convert") as start:
                data = QMimeData()
                data.setUrls([QUrl.fromLocalFile(str(source))])
                drop = QDropEvent(QPointF(5, 5), Qt.CopyAction, data, Qt.LeftButton, Qt.ShiftModifier)
                window.drop.dropEvent(drop)
                self.assertEqual(window.format.currentText(), "docx")
                start.assert_called_once()
            window.close()

    def test_mixed_formats_and_numbered_outputs(self):
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "same.png"
            Image.new("RGB", (20, 20), "red").save(image)
            text = root / "same.txt"
            text.write_text("hello", encoding="utf-8")
            image, text = image.resolve(), text.resolve()
            (root / "same.pdf").write_bytes(b"existing")
            (root / "same.docx").write_bytes(b"existing word file")
            window = Window()
            window.output_dir = root
            window._add_files([image, text])
            self.assertEqual(window.formats[image], "webp")
            window.list.setCurrentRow(1)
            window.format.setCurrentText("docx")
            self.assertEqual(window.formats[text], "docx")
            self.assertEqual(window.formats[image], "webp")
            reserved = {str(image).casefold(), str(text).casefold()}
            first = numbered_target(root, "same", ".pdf", reserved)
            second = numbered_target(root, "same", ".pdf", reserved)
            self.assertEqual((first.name, second.name), ("same (2).pdf", "same (3).pdf"))
            window._convert()
            self.assertTrue(window.worker.wait(10000))
            app.processEvents()
            self.assertTrue((root / "same.webp").exists())
            self.assertEqual(Document(root / "same (2).docx").paragraphs[0].text, "hello")
            self.assertEqual((root / "same.docx").read_bytes(), b"existing word file")
            window.close()

    def test_pdf_tools_keep_pages_and_sources(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            first = root / "first.pdf"
            second = root / "second.pdf"
            (root / "first.txt").write_text("FIRST", encoding="utf-8")
            (root / "second.txt").write_text("SECOND", encoding="utf-8")
            convert(root / "first.txt", first)
            convert(root / "second.txt", second)
            merged = root / "merged.pdf"
            pdf_merge([first, second], merged)
            self.assertEqual(len(PdfReader(str(merged)).pages), 2)
            (root / "merged-2.png").write_bytes(b"existing page image")
            self.assertEqual(numbered_target(root, "merged", ".png", set(), pages=2).name, "merged (2).png")
            reordered = root / "reordered.pdf"
            pdf_reorder(merged, reordered, [2, 1])
            self.assertIn("SECOND", PdfReader(str(reordered)).pages[0].extract_text())
            rotated = root / "rotated.pdf"
            pdf_rotate(first, rotated, 90)
            self.assertEqual(PdfReader(str(rotated)).pages[0].rotation, 90)
            pages = root / "pages"
            pdf_split(merged, pages)
            self.assertEqual(sorted(path.name for path in pages.iterdir()), ["page-1.pdf", "page-2.pdf"])
            with self.assertRaises(ValueError):
                pdf_reorder(merged, root / "bad.pdf", [1, 1])
            self.assertIn("FIRST", PdfReader(str(first)).pages[0].extract_text())

    def test_image_tools_resize_crop_and_strip_metadata(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "photo.png"
            Image.new("RGB", (60, 40), "blue").save(source, pnginfo=self._comment())
            resized = root / "resized.png"
            edit_image(source, resized, "resize", (30, 30))
            with Image.open(resized) as image:
                self.assertEqual(image.size, (30, 20))
            cropped = root / "cropped.png"
            edit_image(source, cropped, "crop", (10, 5, 20, 15))
            with Image.open(cropped) as image:
                self.assertEqual(image.size, (20, 15))
            stripped = root / "stripped.png"
            edit_image(source, stripped, "strip_metadata")
            with Image.open(stripped) as image:
                self.assertNotIn("Comment", image.info)
            with self.assertRaises(ValueError):
                edit_image(source, root / "bad.png", "crop", (50, 0, 20, 20))

    def test_media_trim_keeps_source(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "tone.wav"
            with wave.open(str(source), "wb") as output:
                output.setparams((1, 2, 8000, 0, "NONE", "not compressed"))
                output.writeframes(b"\x00\x00" * 16000)
            trimmed = root / "tone-trimmed.wav"
            trim_media(source, trimmed, 0.5, 1.5)
            with wave.open(str(trimmed), "rb") as result:
                self.assertAlmostEqual(result.getnframes() / result.getframerate(), 1, delta=0.1)
            with self.assertRaises(FileExistsError):
                trim_media(source, trimmed, 0.5, 1.5)
            self.assertTrue(source.exists())

    def test_zip_preview_and_safe_extraction(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            archive = root / "files.zip"
            with ZipFile(archive, "w") as output:
                output.writestr("inside/note.txt", "hello")
            self.assertEqual(zip_entries(archive)[0][0], "inside/note.txt")
            target = root / "extracted"
            extract_zip(archive, target)
            self.assertEqual((target / "inside" / "note.txt").read_text(), "hello")
            with self.assertRaises(FileExistsError):
                extract_zip(archive, target)
            bad = root / "bad.zip"
            with ZipFile(bad, "w") as output:
                output.writestr("../escape.txt", "no")
            with self.assertRaises(ValueError):
                extract_zip(bad, root / "unsafe")
            self.assertFalse((root / "escape.txt").exists())

    def test_ocr_image_and_pdf_pages(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "scan.png"
            scan = Image.new("RGB", (600, 160), "white")
            ImageDraw.Draw(scan).text((25, 35), "HELLO OCR 123", fill="black", font=ImageFont.load_default(size=48))
            scan.save(image)
            ocr_to_text(image, root / "image.txt")
            self.assertIn("HELLO", (root / "image.txt").read_text())
            pdf = root / "pages.pdf"
            convert(image, pdf)
            ocr_to_text(pdf, root / "pages.txt")
            self.assertIn("HELLO", (root / "pages.txt").read_text())
            mixed = root / "mixed.pdf"
            document = canvas.Canvas(str(mixed))
            document.drawString(40, 700, "DIGITAL TEXT")
            document.showPage()
            document.drawImage(str(image), 0, 500, width=600, height=160)
            document.save()
            ocr_to_text(mixed, root / "mixed.txt")
            result = (root / "mixed.txt").read_text()
            self.assertIn("DIGITAL TEXT", result)
            self.assertIn("HELLO", result)

    def test_ocr_multilingual_models(self):
        fixtures = Path(__file__).parent / "test_data"
        samples = (("multilingual", "latin", "CAFÉ"),
                   ("cyrillic", "cyrillic", "Привет"),
                   ("hindi", "hindi", "भारत"),
                   ("korean", "korean", "세계"),
                   ("arabic", "arabic", "مرحبا"))
        with tempfile.TemporaryDirectory() as folder:
            with patch("socket.socket.connect", side_effect=AssertionError("OCR tried to access the network")):
                for language, name, expected in samples:
                    output = Path(folder) / f"{name}.txt"
                    ocr_to_text(fixtures / f"ocr_{name}.png", output, language)
                    self.assertIn(expected, output.read_text(encoding="utf-8"))

    @staticmethod
    def _comment():
        from PIL.PngImagePlugin import PngInfo
        metadata = PngInfo()
        metadata.add_text("Comment", "private")
        return metadata


if __name__ == "__main__":
    unittest.main()
