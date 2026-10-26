"""Native active-sheet print settings without rebuilding workbook parts."""
import math,re,tempfile
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from openpyxl.utils.cell import range_boundaries,get_column_letter,quote_sheetname
from PySide6.QtWidgets import (QDialog,QFormLayout,QLineEdit,QComboBox,QSpinBox,
    QDoubleSpinBox,QCheckBox,QDialogButtonBox,QLabel)
from office_surgery import S,PARSER,_active_sheet
from office_editing import command

PAPERS={'Letter':'1','Legal':'5','A3':'8','A4':'9'}
DEFAULTS=dict(paper='A4',orientation='portrait',mode='Fit',scale=100,wide=1,tall=0,
              area='',repeat='',left=.7,right=.7,top=.75,bottom=.75,gridlines=False,headings=False,center=False)
ORDER='sheetPr dimension sheetViews sheetFormatPr cols sheetData sheetCalcPr sheetProtection protectedRanges scenarios autoFilter sortState dataConsolidate customSheetViews mergeCells phoneticPr conditionalFormatting dataValidations hyperlinks printOptions pageMargins pageSetup headerFooter rowBreaks colBreaks customProperties cellWatches ignoredErrors smartTags drawing legacyDrawing legacyDrawingHF picture oleObjects controls webPublishItems tableParts extLst'.split()


def areas(text):
    if not isinstance(text,str) or len(text)>2000:raise ValueError('Print area is too long.')
    result=[]
    for part in text.split(',') if text.strip() else []:
        part=part.strip().upper()
        if not re.fullmatch(r'\$?[A-Z]{1,3}\$?[1-9]\d*(?::\$?[A-Z]{1,3}\$?[1-9]\d*)?',part):raise ValueError('Use cell ranges such as A1:F40, H1:J20.')
        left,top,right,bottom=range_boundaries(part)
        if not (1<=left<=right<=16384 and 1<=top<=bottom<=1048576):raise ValueError('Print area is outside worksheet bounds or reversed.')
        result.append(f'${get_column_letter(left)}${top}:${get_column_letter(right)}${bottom}')
    return result


def validate(settings):
    if settings.get('paper') not in PAPERS or settings.get('orientation') not in ('portrait','landscape') or settings.get('mode') not in ('Fit','Percentage'):raise ValueError('Choose a supported paper, orientation and scaling mode.')
    for key,low,high in [('scale',10,400),('wide',0,99),('tall',0,99)]:
        if type(settings.get(key)) is not int or not low<=settings[key]<=high:raise ValueError('Use 10–400 percent and 0–99 pages for fit dimensions (0 is automatic).')
    for key in ('left','right','top','bottom'):
        value=settings.get(key)
        if not isinstance(value,(int,float)) or not math.isfinite(value) or not 0<=value<=5:raise ValueError('Margins must be between 0 and 5 inches.')
    width,height={'Letter':(8.5,11),'Legal':(8.5,14),'A3':(11.693,16.535),'A4':(8.268,11.693)}[settings['paper']]
    if settings['orientation']=='landscape':width,height=height,width
    if settings['left']+settings['right']>=width-.1 or settings['top']+settings['bottom']>=height-.1:raise ValueError('Margins must leave a printable area on the selected paper.')
    areas(settings.get('area'))
    repeat=settings.get('repeat')
    if not isinstance(repeat,str) or (repeat and not re.fullmatch(r'[1-9]\d*:[1-9]\d*',repeat)):raise ValueError('Repeat rows using a range such as 1:2, or leave blank.')
    if repeat and not 1<=int(repeat.split(':')[0])<=int(repeat.split(':')[1])<=1048576:raise ValueError('Repeating rows are outside worksheet bounds or reversed.')


def _tokens(value):return re.findall(r"(?:'(?:[^']|'')*'|[^,])+",value or '')


def load_layout(editor,path=None):
    editor.print_layout=dict(DEFAULTS);editor.pending_print_layout=False
    if path is None or Path(path).suffix.lower()!='.xlsx':return
    with ZipFile(path) as archive:
        sheet=etree.fromstring(archive.read(_active_sheet(archive)),PARSER);book=etree.fromstring(archive.read('xl/workbook.xml'),PARSER)
        active=book.find(f'{{{S}}}bookViews/{{{S}}}workbookView');index=int(active.get('activeTab','0')) if active is not None else 0
        config=editor.print_layout;setup=sheet.find(f'{{{S}}}pageSetup');margins=sheet.find(f'{{{S}}}pageMargins');options=sheet.find(f'{{{S}}}printOptions')
        if setup is not None:
            config.update(paper=next((key for key,value in PAPERS.items() if value==setup.get('paperSize')),'A4'),orientation=setup.get('orientation','portrait'),scale=int(setup.get('scale','100')),wide=int(setup.get('fitToWidth','1')),tall=int(setup.get('fitToHeight','0')))
        props=sheet.find(f'{{{S}}}sheetPr/{{{S}}}pageSetUpPr');config['mode']='Fit' if props is not None and props.get('fitToPage') in ('1','true') else 'Percentage'
        if margins is not None:
            for key in ('left','right','top','bottom'):config[key]=float(margins.get(key,str(config[key])))
        if options is not None:
            for key,attr in [('gridlines','gridLines'),('headings','headings'),('center','horizontalCentered')]:config[key]=options.get(attr) in ('1','true')
        for node in book.findall(f'{{{S}}}definedNames/{{{S}}}definedName'):
            if node.get('localSheetId')!=str(index):continue
            refs=[v.rsplit('!',1)[-1] for v in _tokens(node.text)]
            if node.get('name')=='_xlnm.Print_Area':config['area']=','.join(refs)
            elif node.get('name')=='_xlnm.Print_Titles':config['repeat']=next((r.replace('$','') for r in refs if re.fullmatch(r'\$?\d+:\$?\d+',r)),'')


def set_layout(editor,settings):
    validate(settings);editor.print_layout=dict(settings);editor.pending_print_layout=True;editor.dirty=True
    editor.statusBar().showMessage('Print settings apply to the active worksheet when saved as XLSX. The grid view is unchanged.')


def save_layout(editor,path):
    if not editor.pending_print_layout:return
    settings=editor.print_layout;validate(settings);tag=lambda name:'{'+S+'}'+name
    with ZipFile(path) as archive:
        if len(archive.namelist())!=len(set(archive.namelist())) or any(n.startswith('_xmlsignatures/') for n in archive.namelist()):raise ValueError('Signed or duplicate-part workbooks cannot be edited safely.')
        sheet_path=_active_sheet(archive);sheet=etree.fromstring(archive.read(sheet_path),PARSER);book=etree.fromstring(archive.read('xl/workbook.xml'),PARSER)
        def ensure(parent,name):
            node=parent.find(tag(name))
            if node is None:
                node=etree.Element(tag(name));rank=ORDER.index(name)
                following=next((n for n in parent if etree.QName(n).localname in ORDER and ORDER.index(etree.QName(n).localname)>rank),None)
                parent.insert(parent.index(following) if following is not None else len(parent),node)
            return node
        props=ensure(sheet,'sheetPr');fit=props.find(tag('pageSetUpPr'))
        if fit is None:fit=etree.SubElement(props,tag('pageSetUpPr'))
        fit.set('fitToPage','1' if settings['mode']=='Fit' else '0')
        setup=ensure(sheet,'pageSetup');setup.set('paperSize',PAPERS[settings['paper']]);setup.set('orientation',settings['orientation'])
        for key,attr in [('scale','scale'),('wide','fitToWidth'),('tall','fitToHeight')]:setup.set(attr,str(settings[key]))
        margins=ensure(sheet,'pageMargins')
        for key in ('left','right','top','bottom'):margins.set(key,str(settings[key]))
        for key in ('header','footer'):
            if margins.get(key) is None:margins.set(key,'.3')
        options=ensure(sheet,'printOptions')
        for key,attr in [('gridlines','gridLines'),('headings','headings'),('center','horizontalCentered')]:options.set(attr,'1' if settings[key] else '0')
        view=book.find(tag('bookViews')+'/'+tag('workbookView'));index=int(view.get('activeTab','0')) if view is not None else 0
        name=quote_sheetname(list(book.find(tag('sheets')))[index].get('name'))
        names=book.find(tag('definedNames'))
        if names is None:
            names=etree.Element(tag('definedNames'));following=next((n for n in book if etree.QName(n).localname in ('calcPr','oleSize','customWorkbookViews','pivotCaches','smartTagPr','smartTagTypes','webPublishing','fileRecoveryPr','webPublishObjects','extLst')),None);book.insert(book.index(following) if following is not None else len(book),names)
        def definition(kind,values):
            nodes=[n for n in names if n.get('name')==kind and n.get('localSheetId')==str(index)]
            for node in nodes:names.remove(node)
            if values:etree.SubElement(names,tag('definedName'),name=kind,localSheetId=str(index)).text=','.join(values)
        titles=next((n.text for n in names if n.get('name')=='_xlnm.Print_Titles' and n.get('localSheetId')==str(index)),'')
        keep=[v for v in _tokens(titles) if not re.fullmatch(r'\$?\d+:\$?\d+',v.rsplit('!',1)[-1])]
        if settings['repeat']:
            first,last=settings['repeat'].split(':');keep.append(f'{name}!${first}:${last}')
        definition('_xlnm.Print_Titles',keep);definition('_xlnm.Print_Area',[f'{name}!{v}' for v in areas(settings['area'])])
        replacements={sheet_path:etree.tostring(sheet,xml_declaration=True,encoding='UTF-8'),'xl/workbook.xml':etree.tostring(book,xml_declaration=True,encoding='UTF-8')}
        with tempfile.NamedTemporaryFile(dir=Path(path).parent,suffix='.xlsx',delete=False) as stream:temporary=Path(stream.name)
        try:
            with ZipFile(temporary,'w') as output:
                for info in archive.infolist():output.writestr(info,replacements.get(info.filename,archive.read(info.filename)))
        except Exception:temporary.unlink(missing_ok=True);raise
    try:temporary.replace(path)
    finally:temporary.unlink(missing_ok=True)
    editor.pending_print_layout=False


def edit_layout(editor):
    dialog=QDialog(editor);dialog.setWindowTitle('Worksheet print layout');form=QFormLayout(dialog);widgets={}
    form.addRow(QLabel('Native XLSX print settings for the active worksheet.\nThe grid is not a print preview. Excel/LibreOffice PDF conversion uses these settings.'))
    for key,label,values in [('paper','Paper',list(PAPERS)),('orientation','Orientation',['portrait','landscape']),('mode','Scaling mode',['Fit','Percentage'])]:
        field=QComboBox();field.addItems(values);field.setCurrentText(editor.print_layout[key]);widgets[key]=field;form.addRow(label,field)
    for key,label,low,high in [('scale','Scale (%)',10,400),('wide','Fit pages wide (0 = auto)',0,99),('tall','Fit pages tall (0 = auto)',0,99)]:
        field=QSpinBox();field.setRange(low,high);field.setValue(editor.print_layout[key]);widgets[key]=field;form.addRow(label,field)
    for key,label in [('area','Print area (blank = all used cells)'),('repeat','Repeat rows (e.g. 1:2)')]:
        field=QLineEdit(editor.print_layout[key]);widgets[key]=field;form.addRow(label,field)
    for key in ('left','right','top','bottom'):
        field=QDoubleSpinBox();field.setRange(0,5);field.setDecimals(3);field.setSuffix(' in');field.setValue(editor.print_layout[key]);widgets[key]=field;form.addRow(key.title()+' margin',field)
    for key,label in [('gridlines','Print gridlines'),('headings','Print row/column headings'),('center','Center horizontally')]:
        field=QCheckBox();field.setChecked(editor.print_layout[key]);widgets[key]=field;form.addRow(label,field)
    for key,field in widgets.items():field.setAccessibleName('Print '+key)
    def scaling_mode():
        fit=widgets['mode'].currentText()=='Fit';widgets['scale'].setEnabled(not fit)
        for key in ('wide','tall'):widgets[key].setEnabled(fit)
    widgets['mode'].currentTextChanged.connect(scaling_mode);scaling_mode()
    buttons=QDialogButtonBox(QDialogButtonBox.Ok|QDialogButtonBox.Cancel);buttons.accepted.connect(dialog.accept);buttons.rejected.connect(dialog.reject);form.addRow(buttons)
    if dialog.exec()==QDialog.Accepted:
        settings={key:field.currentText() if isinstance(field,QComboBox) else field.isChecked() if isinstance(field,QCheckBox) else field.text().strip() if isinstance(field,QLineEdit) else field.value() for key,field in widgets.items()};set_layout(editor,settings)


def setup_print_tools(editor):
    load_layout(editor)
    command(editor,'Layout','Worksheet print layout',lambda:edit_layout(editor),'settings')
