import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import tempfile,unittest
from pathlib import Path
from zipfile import ZipFile
from openpyxl import Workbook,load_workbook
from openpyxl.chart import BarChart,Reference
from PySide6.QtWidgets import QApplication
from PySide6.QtCore import QCoreApplication,QEvent
from office_suite import Tables
from office_print_layout import set_layout,areas,DEFAULTS
from office_ribbon import install_ribbon


class PrintLayoutTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def tearDown(self):
        for widget in self.app.topLevelWidgets():widget.deleteLater()
        QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete)

    def test_native_settings_reload_and_unrelated_parts_survive(self):
        with tempfile.TemporaryDirectory() as folder:
            source=Path(folder)/'original.xlsx';target=Path(folder)/'print.xlsx'
            book=Workbook();sheet=book.active;sheet.title="Report, team's";sheet.append(['Item','Count']);sheet.append(['Local',10]);sheet['C2']='=B2*2';sheet.print_title_cols='A:B';other=book.create_sheet('Untouched');other['A1']='Keep me'
            chart=BarChart();chart.add_data(Reference(sheet,min_col=2,min_row=1,max_row=2),titles_from_data=True);sheet.add_chart(chart,'F3');book.save(source);book.close()
            editor=Tables();editor._read(source);set_layout(editor,dict(DEFAULTS,paper='Letter',orientation='landscape',area='A1:C30, E1:G10',repeat='1:2',gridlines=True,headings=True,center=True,left=.4))
            with self.assertRaises(ValueError):editor._write(source)
            with self.assertRaises(ValueError):editor._write(Path(folder)/'wrong.csv')
            editor._write(target);self.assertFalse(editor.pending_print_layout)
            with ZipFile(source) as old,ZipFile(target) as new:
                changed=[name for name in old.namelist() if old.read(name)!=new.read(name)]
                self.assertEqual(set(changed),{'xl/workbook.xml','xl/worksheets/sheet1.xml'})
            result=load_workbook(target);sheet=result.active
            self.assertEqual(sheet.page_setup.orientation,'landscape');self.assertEqual(sheet.page_setup.paperSize,1)
            self.assertEqual(sheet.print_title_rows,'$1:$2');self.assertEqual(sheet.print_title_cols,'$A:$B')
            self.assertTrue(sheet.print_options.gridLines);self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
            self.assertEqual(sheet['C2'].value,'=B2*2');self.assertEqual(len(sheet._charts),1);self.assertEqual(result['Untouched']['A1'].value,'Keep me');result.close()
            editor._read(target);self.assertEqual(editor.print_layout['repeat'],'1:2');self.assertEqual(len(areas(editor.print_layout['area'])),2)
            set_layout(editor,dict(editor.print_layout,repeat='',area='',mode='Percentage',scale=80));editor._write(Path(folder)/'again.xlsx')
            result=load_workbook(Path(folder)/'again.xlsx');self.assertIsNone(result.active.print_title_rows);self.assertEqual(result.active.print_title_cols,'$A:$B');self.assertFalse(result.active.sheet_properties.pageSetUpPr.fitToPage);result.close()

    def test_validation_new_workbook_and_ribbon(self):
        editor=Tables();install_ribbon(editor);self.assertIn('Worksheet print layout',editor.ribbon_buttons)
        for area in ('XFE1:XFE2','A0:B2','B2:A1','A1;B2','1:3'):
            with self.assertRaises(ValueError):areas(area)
        for settings in (dict(DEFAULTS,repeat='4:2'),dict(DEFAULTS,left=float('nan')),dict(DEFAULTS,wide=-1),dict(DEFAULTS,left=5,right=5)):
            with self.assertRaises(ValueError):set_layout(editor,settings)
        with tempfile.TemporaryDirectory() as folder:
            set_layout(editor,dict(DEFAULTS));path=Path(folder)/'new.xlsx';editor._write(path)
            result=load_workbook(path);self.assertEqual(result.active.page_setup.paperSize,9);result.close()
        editor.new();self.assertFalse(editor.pending_print_layout)


if __name__=='__main__':unittest.main()
