"""Editing batches must remain usable and survive native-format round trips."""
import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import tempfile,unittest
from pathlib import Path
from zipfile import ZipFile
from unittest.mock import patch
from PySide6.QtCore import Qt,QPointF,QEvent,QCoreApplication
from PySide6.QtGui import QEnterEvent,QTextCursor,QTextCharFormat,QColor
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication,QTableWidgetItem
from docx import Document
from openpyxl import Workbook,load_workbook
from pptx import Presentation
from office_suite import Pages,Tables,Dazzle,Sketch,Organize
from office_ribbon import install_ribbon
from office_editing import heading,page_setup,paragraph_spacing,replace_text,duplicate_slide,align_object
from office_cells import format_cells,paste_cells


class RibbonBatchTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def tearDown(self):
        for widget in self.app.topLevelWidgets():widget.deleteLater()
        QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete)

    def test_drop_hover_resets_and_click_keyboard_choose_once(self):
        from app import DropArea
        drop=DropArea();chosen=[];drop.choose_requested.connect(lambda:chosen.append(True));drop.show()
        drop.enterEvent(QEnterEvent(QPointF(20,20),QPointF(20,20),QPointF(20,20)));QTest.qWait(220);self.assertGreater(drop._drop_level,.95)
        drop.leaveEvent(QEvent(QEvent.Leave));QTest.qWait(220);self.assertEqual(drop._drop_level,0)
        QTest.mouseClick(drop,Qt.LeftButton);QTest.keyClick(drop,Qt.Key_Space);self.assertEqual(len(chosen),2)
        drop.hide();self.assertEqual(drop._drop_level,0)

    def test_ribbon_preserves_actions_widgets_shortcuts_and_small_layout(self):
        from formulaic import Formulaic
        for cls in (Pages,Tables,Dazzle,Sketch,Organize,Formulaic):
            window=cls();window.resize(780,580);window.show();self.app.processEvents()
            self.assertIn('Home',[window.ribbon.tabText(i) for i in range(window.ribbon.count())]);open_label='Open database' if isinstance(window,Organize) else 'Open';self.assertIn(open_label,window.ribbon_buttons)
            if 'Save' in window.ribbon_buttons:self.assertEqual(window.ribbon_buttons['Save'].defaultAction().shortcut().toString(),'Ctrl+S')
            with patch('PySide6.QtWidgets.QFileDialog.getOpenFileName',return_value=('','')) as picker:
                window.ribbon_buttons[open_label].click();self.assertTrue(picker.called)
            if isinstance(window,Pages):
                self.assertTrue(window.font_family.isVisible());self.assertTrue(window.font_size.isVisible());self.assertLess(window.font_size.geometry().bottom(),window.ribbon.height())
            window.hide()

    def test_pages_native_styles_spacing_layout_and_find_undo(self):
        with tempfile.TemporaryDirectory() as temporary:
            path=Path(temporary)/'styled.docx';pages=Pages();pages.editor.setPlainText('Title\nAlpha alpha alphabet\nLast')
            cursor=pages.editor.textCursor();cursor.setPosition(0);pages.editor.setTextCursor(cursor);heading(pages,1);paragraph_spacing(pages,1.5)
            fmt=QTextCharFormat();fmt.setFontStrikeOut(True);fmt.setBackground(QColor('#ffdd22'));pages.editor.mergeCurrentCharFormat(fmt)
            page_setup(pages,dict(paper='A4',landscape=True,margin=17))
            count=replace_text(pages.editor,'alpha','Beta',False,True);self.assertEqual(count,2);self.assertIn('alphabet',pages.editor.toPlainText());pages.editor.undo();self.assertIn('Alpha alpha',pages.editor.toPlainText())
            pages._write(path);doc=Document(path);self.assertEqual(doc.paragraphs[0].style.name,'Heading 1');self.assertAlmostEqual(doc.paragraphs[0].paragraph_format.line_spacing,1.5);self.assertGreater(doc.sections[0].page_width,doc.sections[0].page_height);self.assertAlmostEqual(doc.sections[0].left_margin.mm,17,places=1)
            other=Pages();other.open(path);self.assertEqual(other.editor.document().firstBlock().blockFormat().headingLevel(),1)

    def test_xlsx_formatting_blank_cells_and_unrelated_parts_roundtrip(self):
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary);source=root/'source.xlsx';target=root/'result.xlsx';second=root/'result2.xlsx';book=Workbook();book.active.append(['Value',12]);book.create_sheet('Unrelated').append(['Preserve me']);book.save(source);book.close()
            editor=Tables();editor.open(source);editor.grid.setCurrentCell(0,1);format_cells(editor,name='Arial',size=18,bold=True,fill='#ffdd22',color='#111111',horizontal='center',number_format='0.00%');editor.grid.setCurrentCell(7,4);format_cells(editor,italic=True);editor._write(target)
            result=load_workbook(target);cell=result.active['B1'];self.assertEqual(cell.font.name,'Arial');self.assertTrue(cell.font.bold);self.assertEqual(cell.fill.fgColor.rgb,'FFFFDD22');self.assertEqual(cell.number_format,'0.00%');self.assertTrue(result.active['E8'].font.italic);result.close()
            with ZipFile(source) as a,ZipFile(target) as b:self.assertEqual(a.read('xl/worksheets/sheet2.xml'),b.read('xl/worksheets/sheet2.xml'))
            reopened=Tables();reopened.open(target);self.assertTrue(reopened.grid.item(0,1).font().bold());reopened.grid.setCurrentCell(0,1);format_cells(reopened,italic=True);reopened._write(second)
            result=load_workbook(second);self.assertTrue(result.active['B1'].font.bold and result.active['B1'].font.italic);result.close()
            reopened.grid.setCurrentCell(3,0);paste_cells(reopened,'one\ttwo\nthree\tfour');self.assertEqual(reopened.grid.item(4,1).text(),'four')

    def test_duplicate_slide_independent_chart_media_timing_and_alignment(self):
        from office_insert import chart_data
        from dazzle import native_chart_data,effects
        with tempfile.TemporaryDirectory() as temporary:
            editor=Dazzle();editor.insert_text('Keep original');chart=editor.insert_chart('Column',chart_data('X,Y\nA,2\nB,3'));editor.animate_object(chart,.5);duplicate_slide(editor)
            self.assertEqual(len(editor.deck.slides),2);chart_copy=next(s for s in editor.deck.slides[1].shapes if s.has_chart);chart_copy.chart.replace_data(native_chart_data('Column',chart_data('X,Y\nA,8\nB,9')))
            self.assertEqual(list(editor.deck.slides[0].shapes[-1].chart.series[0].values),[2.,3.]);editor.choose(chart_copy);align_object(editor,'Right');self.assertEqual(editor.selected_shape().left+editor.selected_shape().width,editor.deck.slide_width)
            editor.animate_object(editor.selected_shape(),None);self.assertFalse(effects(editor.deck.slides[1]));path=Path(temporary)/'copy.pptx';editor._write(path);deck=Presentation(path);self.assertEqual(list(next(s for s in deck.slides[1].shapes if s.has_chart).chart.series[0].values),[8.,9.])


if __name__=='__main__':unittest.main()
